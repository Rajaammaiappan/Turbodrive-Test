# -*- coding: utf-8 -*-
"""
STELOS - The ASD-STE100 Intelligence Platform (an ALTEN product).
Simplified Technical English (ASD-STE100 Issue 9) checker
with in-place word replacement into a COPY of the source file.
RR / ALTEN internal tool. Bind 127.0.0.1, port 5002.

Inputs supported: pasted text, .docx, .txt/.md, .xlsx/.xlsm
Replacement:  tick the words to fix -> tool writes a corrected COPY and offers it
              for download. The original uploaded file is NEVER modified.

Ships with ste_dictionary.json (ASD-STE100 Issue 9, 2025-01-15) in the SAME folder.
"""
import os, re, json, csv, sqlite3, threading, webbrowser, time, traceback, html, uuid, tempfile
from datetime import datetime
from flask import Flask, request, jsonify, render_template_string, send_file
from werkzeug.exceptions import RequestEntityTooLarge

app = Flask(__name__)
PORT = 5002
HERE = os.path.dirname(os.path.abspath(__file__))
DATA_PATH = os.path.join(HERE, 'ste_dictionary.json')
TMPDIR = os.path.join(tempfile.gettempdir(), 'ste_checker')
os.makedirs(TMPDIR, exist_ok=True)

# Flask has no upload size cap by default, which risks an unbounded upload
# hanging the server / exhausting memory. Cap it generously (engineering
# docs with embedded images can be large) and return a JSON error the UI
# can show, instead of Flask's default HTML error page.
MAX_UPLOAD_MB = int(os.environ.get('STE_MAX_UPLOAD_MB', '200'))
app.config['MAX_CONTENT_LENGTH'] = MAX_UPLOAD_MB * 1024 * 1024

@app.errorhandler(RequestEntityTooLarge)
def _too_large(_e):
    return jsonify({'error': 'File is too large (max %d MB). Split the document or '
                              'raise STE_MAX_UPLOAD_MB.' % MAX_UPLOAD_MB}), 413

# ---------------------------------------------------------------- load data
with open(DATA_PATH, 'r', encoding='utf-8') as f:
    STE = json.load(f)
APPROVED = STE['approved']
NOT_APPROVED = STE['not_approved']
RULES = STE['rules']
META = STE['meta']

ALWAYS_OK = set("a an the and or of to in on at for from with by as is are was were "
                "be do this that it not no if".split())

# ------------------------------------------------------------------ item 3
# ASD-STE100 issue/version shown in the UI and reports. Read from the
# dictionary meta if present, else a sane default; STE_DOC_URL lets the
# backend point at the controlling specification document (a SharePoint /
# intranet link to the official ASD-STE100 Issue 9 PDF, for example).
STE_ISSUE = str(META.get('issue') or '9')
STE_ISSUE_DATE = str(META.get('issue_date') or '2025-01-15')
STE_SOURCE = str(META.get('source') or ('ASD-STE100 Issue %s' % STE_ISSUE))
STE_DOC_URL = os.environ.get('STE_DOC_URL', '').strip()

# ------------------------------------------------------------------ item 8
# Human-readable part-of-speech names, for showing approved usage such as
# "N", "V", "ADJ" as "noun", "verb", "adjective" in guidance output.
POS_FULL = {
 'n': 'noun', 'v': 'verb', 'adj': 'adjective', 'adv': 'adverb',
 'prep': 'preposition', 'conj': 'conjunction', 'pro': 'pronoun',
 'art': 'article', 'det': 'determiner', 'int': 'interjection',
 'tn': 'technical noun', 'tv': 'technical verb',
}
def pos_label(pos):
    """'v' -> 'V (verb)', 'n/v' -> 'N (noun), V (verb)'. Blank -> ''."""
    if not pos:
        return ''
    parts = re.split(r'[\s/,]+', str(pos).strip())
    out = []
    for p in parts:
        pl = p.lower().strip('.')
        if not pl:
            continue
        full = POS_FULL.get(pl)
        out.append('%s (%s)' % (p.upper(), full) if full else p.upper())
    return ', '.join(out)

def approved_pos(word):
    """Return the approved part(s) of speech recorded for an approved word
    (may be '' when the dictionary did not record one)."""
    for f in deinflect(word):
        e = APPROVED.get(f)
        if e and e.get('pos'):
            return e['pos']
    return ''

# ------------------------------------------------------------------ item 1
# Repair words that arrive split by stray spaces / invisible characters
# (common when text is copied out of a PDF or a DOCX with optional hyphens):
# "use d" -> "used", "mix ed" -> "mixed", "con- \n nect" -> "connect".
# Only rejoins when the joined form is a real dictionary word, so it never
# merges two legitimate separate words.
INVISIBLE_RE = re.compile('[­​‌‍﻿]')
DEHYPHEN_RE = re.compile(r'([A-Za-z])[­-]\s*\n\s*([A-Za-z])')
SPLIT_SUFFIX_RE = re.compile(
    r"\b([A-Za-z]{2,})[ \t]+(d|ed|s|es|ing|ings|ly|er|ers|ed)\b", re.IGNORECASE)

_EXTRA_KNOWN = {'key': None, 'words': frozenset()}

def _extra_known_words():
    """Words the tool has learned from the shared technical library and from
    documents uploaded so far (corpus words seen at least twice). Used only to
    decide whether a repaired word is a real word. Re-read when the files change."""
    def _mt(p):
        try:
            return os.path.getmtime(p)
        except OSError:
            return 0
    key = (_mt(TERMS_FILE), _mt(CORPUS_FILE))
    if _EXTRA_KNOWN['key'] != key:
        words = set(load_terms().keys())
        try:
            with open(CORPUS_FILE, 'r', encoding='utf-8') as fh:
                words |= {w for w, c in json.load(fh).items() if c >= 2}
        except (FileNotFoundError, ValueError, OSError):
            pass
        _EXTRA_KNOWN['key'] = key
        _EXTRA_KNOWN['words'] = frozenset(words)
    return _EXTRA_KNOWN['words']

def _known_word(w):
    wl = w.lower()
    if wl in APPROVED or wl in NOT_APPROVED:
        return True
    if any(f in APPROVED or f in NOT_APPROVED for f in deinflect(wl)):
        return True
    extra = _extra_known_words()
    return wl in extra or any(f in extra for f in deinflect(wl))

def normalize_text(text, notes=None):
    """Clean invisible characters and repair line-break hyphenation and
    stray-space word splits. Conservative: a split is only rejoined when the
    rejoined token is a recognized word, so real words are never merged.
    When a split looks broken ("xyz ed") but the joined word is NOT known, the
    text is left alone and a note is added to `notes` for the author to check."""
    if not text:
        return text
    text = INVISIBLE_RE.sub('', text)
    text = DEHYPHEN_RE.sub(lambda m: m.group(1) + m.group(2), text)
    def _join(m):
        stem, suf = m.group(1), m.group(2)
        joined = stem + suf
        # only rejoin when the split fragment is NOT itself a real word and
        # the joined form IS a real word (avoids merging "the s", "use it")
        if _known_word(joined) and not _known_word(suf):
            return joined
        if (notes is not None and suf.lower() in ('d', 'ed') and len(stem) >= 3
                and stem.lower() not in ALWAYS_OK and len(notes) < 50
                and m.group(0) not in notes):
            notes.append(m.group(0))
        return m.group(0)
    return SPLIT_SUFFIX_RE.sub(_join, text)

# token registry: token -> {path, filename, ext}
UPLOADS = {}

# ---- shared technical-terms library ("DB" file, safe for a network share) ----
# Point STE_TERMS_FILE at a shared drive so all users read/write the same list.
TERMS_FILE = os.environ.get('STE_TERMS_FILE', os.path.join(HERE, 'tech_terms.json'))
TERMS_LOCK = threading.Lock()
# item 5 - separate categories for maintenance / auditing of the shared library
TERM_CATEGORIES = ['Technical Terms', 'Material Names', 'Part Names', 'Miscellaneous']
DEFAULT_CATEGORY = 'Technical Terms'
CURRENT_USER = os.environ.get('USERNAME') or os.environ.get('USER') or 'user'
CURRENT_MACHINE = os.environ.get('COMPUTERNAME') or os.environ.get('HOSTNAME') or 'machine'

# ---- pilot / library-building mode --------------------------------------------
# While the library is being built by a small group of Authors and CAs, the tool
# must not be used to produce final documents. With STE_PILOT_MODE=1
# the "apply replacements" and "annotated document" outputs are switched off and
# the screen says so. Default is OFF (outputs enabled); set STE_PILOT_MODE=1 for pilot mode.
PILOT_MODE = os.environ.get('STE_PILOT_MODE', '0').strip().lower() not in ('0', 'false', 'no', 'off')
PILOT_MSG = ('Pilot mode: STELOS is being used to build the word library. Corrected and annotated '
             'documents are switched off - do not use results as final output.')

# ---- words seen in uploaded documents (feeds word-break repair + library building)
CORPUS_FILE = os.environ.get('STE_CORPUS_FILE', os.path.join(HERE, 'corpus_words.json'))
CORPUS_LOCK = threading.Lock()

def update_corpus(words):
    """Count every alphabetic word (3+ letters) seen in uploaded documents."""
    try:
        counts = {}
        for w in words:
            wl = w.lower().strip("'-")
            if len(wl) >= 3 and wl.isalpha():
                counts[wl] = counts.get(wl, 0) + 1
        if not counts:
            return
        with CORPUS_LOCK:
            try:
                with open(CORPUS_FILE, 'r', encoding='utf-8') as fh:
                    d = json.load(fh)
                if not isinstance(d, dict):
                    d = {}
            except (FileNotFoundError, ValueError, OSError):
                d = {}
            for w, c in counts.items():
                d[w] = d.get(w, 0) + (1 if c else 0)   # +1 per document, not per mention
            tmp = CORPUS_FILE + '.tmp'
            with open(tmp, 'w', encoding='utf-8') as fh:
                json.dump(d, fh, ensure_ascii=False)
            os.replace(tmp, CORPUS_FILE)
    except Exception:
        traceback.print_exc()

# ---- audit trail for the shared library (who added / moved / removed what) -----
HISTORY_FILE = os.environ.get('STE_TERMS_HISTORY', TERMS_FILE + '.history.csv')
HISTORY_FIELDS = ['Timestamp', 'User', 'Action', 'Word', 'From category', 'To category']

def log_term_event(action, word, frm='', to=''):
    try:
        new = not os.path.exists(HISTORY_FILE)
        with open(HISTORY_FILE, 'a', newline='', encoding='utf-8') as fh:
            w = csv.writer(fh)
            if new:
                w.writerow(HISTORY_FIELDS)
            w.writerow([datetime.now().strftime('%Y-%m-%d %H:%M:%S'), CURRENT_USER,
                        action, word, frm, to])
    except Exception:
        traceback.print_exc()

def load_term_history(limit=200):
    try:
        with open(HISTORY_FILE, 'r', newline='', encoding='utf-8') as fh:
            rows = list(csv.DictReader(fh))
        return rows[-limit:][::-1]
    except (FileNotFoundError, OSError):
        return []

# ---- Parts of Speech (POS) reference - editable by the Linguistic Team -------
# pos_definitions.json is created on first start with DRAFT definitions based on
# ASD-STE100 Issue 9 (Rule 1.2 and the dictionary introduction). The Linguistic
# Team can edit that file (no code change) and the appendix updates on reload.
POS_FILE = os.environ.get('STE_POS_FILE', os.path.join(HERE, 'pos_definitions.json'))
POS_DEFAULT = {
 'status': 'Draft - based on ASD-STE100 Issue 9. To be reviewed by the Linguistic Team.',
 'intro': ('Parts of Speech (POS) are the word classes of English: noun, verb, adjective and so on. '
           'STE approves each word for ONE part of speech (Rule 1.2), so you must know the part of '
           'speech to use a word correctly. The position of a word in the sentence shows its part '
           'of speech. Example: "Test" is an approved noun but not an approved verb.'),
 'entries': [
  {'abbr': 'n', 'name': 'Noun', 'definition': 'A word that names a thing, a place, a person or an idea.',
   'ste_tip': 'Usually follows "the", "a" or "an". A noun is not a verb: write "Do a test of the system", not "Test the system" (Rule 1.2, 3.7).',
   'ste': 'Do the leak test of the system.', 'non_ste': 'Test the system for leaks.'},
  {'abbr': 'v', 'name': 'Verb', 'definition': 'A word that shows an action or a state.',
   'ste_tip': 'Use an approved verb to show an action, not a noun (Rule 3.7). In a step, start with the verb: "Remove the bolt."',
   'ste': 'Clean the inner surface of the container.', 'non_ste': 'Perform cleaning of the inner surface.'},
  {'abbr': 'adj', 'name': 'Adjective', 'definition': 'A word that tells you more about a noun (what kind, which one).',
   'ste_tip': 'Comes before a noun or after "is/are". "Dim" is an approved adjective but not an approved verb.',
   'ste': 'A dim light comes on.', 'non_ste': 'Dim the lights.'},
  {'abbr': 'adv', 'name': 'Adverb', 'definition': 'A word that tells you more about a verb, an adjective or another adverb (how, when, where).',
   'ste_tip': 'Use the approved adverb in the dictionary. Do not use an adverb to replace a precise value.',
   'ste': 'Turn the knob slowly.', 'non_ste': ''},
  {'abbr': 'prep', 'name': 'Preposition', 'definition': 'A word that links a noun or pronoun to another word and shows place, time or direction.',
   'ste_tip': 'Approved prepositions are limited. Do not use phrasal verbs (verb + preposition) unless they are in the dictionary (Rule 9.3).',
   'ste': 'Put the cover on the unit.', 'non_ste': ''},
  {'abbr': 'conj', 'name': 'Conjunction', 'definition': 'A word that joins words, groups of words or sentences (and, or, but, if).',
   'ste_tip': 'Use only the approved conjunctions. Do not join two long ideas in one sentence - make two sentences.',
   'ste': 'Open the door and remove the cover.', 'non_ste': ''},
  {'abbr': 'art', 'name': 'Article', 'definition': 'A word that comes before a noun (a, an, the).',
   'ste_tip': 'Use an article to make the noun clear. Use "the" for a specific item.',
   'ste': 'Install the new seal.', 'non_ste': ''},
  {'abbr': 'pro', 'name': 'Pronoun', 'definition': 'A word used in place of a noun (it, they, this).',
   'ste_tip': 'Make sure it is clear which noun the pronoun replaces. If it is not clear, repeat the noun.',
   'ste': '', 'non_ste': ''},
  {'abbr': 'tn / tv', 'name': 'Technical noun / technical verb', 'definition': 'A noun or verb for a product, a process or a subject field that is not in the dictionary. Written by the author, not approved by the dictionary.',
   'ste_tip': 'Allowed under Rules 1.5 and 1.6 if it follows the category rules. In STELOS, add it to the shared library under a category.',
   'ste': '', 'non_ste': ''},
 ]}

def load_pos_ref():
    try:
        with open(POS_FILE, 'r', encoding='utf-8') as fh:
            d = json.load(fh)
        if isinstance(d, dict) and d.get('entries'):
            return d
    except (FileNotFoundError, ValueError, OSError):
        pass
    try:
        with open(POS_FILE, 'w', encoding='utf-8') as fh:
            json.dump(POS_DEFAULT, fh, indent=2, ensure_ascii=False)
    except OSError:
        pass
    return POS_DEFAULT

# ---- sentence-context rules - editable by Authors / CAs ------------------------
# The tool has no language model, so "context" means fixed, visible rules:
#  * cues: the word BEFORE a flagged word shows its job in the sentence
#    (after "the" -> noun, after "to" -> verb ...);
#  * rules: author-added exceptions, e.g. {"word":"set","prev":["the"],"note":"noun"}.
# Context only CLEARS a flag when the word is APPROVED in the role the cue
# shows (e.g. "test" as a noun). Every cleared word is listed on screen.
CONTEXT_FILE = os.environ.get('STE_CONTEXT_FILE', os.path.join(HERE, 'context_rules.json'))
CONTEXT_DEFAULT = {
 'cues': {
  'verb_before': ['to', 'will', 'can', 'must', 'shall', 'should', 'may', 'might', 'could', 'would',
                  'not', 'do', 'does', 'did', 'then', 'and then'],
  'noun_before': ['the', 'a', 'an', 'this', 'that', 'these', 'those', 'each', 'every', 'any',
                  'no', 'some', 'its', 'their', 'your', 'our', 'his', 'her', 'my', 'of', 'another'],
  'adj_before':  ['is', 'are', 'was', 'were', 'be', 'been', 'very', 'too', 'more', 'most'],
 },
 'rules': [],
 'help': ('rules: add {"id":"C1","word":"set","prev":["the","a"],"next":[],"note":"noun"} to stop '
          'flagging a word when the word before (prev) or after (next) is in the list. '
          'Leave a list empty to ignore it.'),
}
_CTX = {'mt': None, 'data': CONTEXT_DEFAULT}

def load_context_rules():
    try:
        mt = os.path.getmtime(CONTEXT_FILE)
    except OSError:
        mt = None
        try:
            with open(CONTEXT_FILE, 'w', encoding='utf-8') as fh:
                json.dump(CONTEXT_DEFAULT, fh, indent=2, ensure_ascii=False)
        except OSError:
            pass
    if mt != _CTX['mt']:
        try:
            with open(CONTEXT_FILE, 'r', encoding='utf-8') as fh:
                d = json.load(fh)
            cues = dict(CONTEXT_DEFAULT['cues'])
            cues.update(d.get('cues') or {})
            _CTX['data'] = {'cues': cues, 'rules': d.get('rules') or []}
        except (FileNotFoundError, ValueError, OSError):
            _CTX['data'] = CONTEXT_DEFAULT
        _CTX['mt'] = mt
    return _CTX['data']

# ---- usage log (SQLite, same shared-log pattern as the manual-comparison
# tool's stelos_log.db - tracks who checked/downloaded what and how many
# documents, for usage/adoption tracking, not for the STE rules themselves) ----
LOG_DB_FILE = os.environ.get('STE_LOG_DB', os.path.join(HERE, 'stelos_log.db'))
LOG_LOCK = threading.Lock()

def _log_db():
    conn = sqlite3.connect(LOG_DB_FILE, timeout=10)
    conn.execute('''CREATE TABLE IF NOT EXISTS usage_log(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            timestamp TEXT, doc_type TEXT, action TEXT, user TEXT, machine TEXT,
            task_no TEXT, task_title TEXT, documents TEXT, doc_count INTEGER,
            mode TEXT, words INTEGER, sentences INTEGER, unapproved INTEGER,
            rule_findings INTEGER, replacements INTEGER, annotations INTEGER,
            rows INTEGER, note TEXT)''')
    return conn

def log_usage(action, documents='', doc_count=1, mode='', words=0, sentences=0,
              unapproved=0, rule_findings=0, replacements=0, annotations=0, note=''):
    """Append one row to the shared usage log. Logging failures never break
    the actual check/download - they're only recorded on a best-effort basis."""
    try:
        with LOG_LOCK:
            conn = _log_db()
            conn.execute(
                '''INSERT INTO usage_log
                   (timestamp, doc_type, action, user, machine, task_no, task_title,
                    documents, doc_count, mode, words, sentences, unapproved,
                    rule_findings, replacements, annotations, rows, note)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (datetime.now().strftime('%Y-%m-%d %H:%M:%S'), 'STE Check', action,
                 CURRENT_USER, CURRENT_MACHINE, '', '', documents, doc_count, mode,
                 words, sentences, unapproved, rule_findings, replacements,
                 annotations, 0, note))
            conn.commit()
            conn.close()
    except Exception:
        traceback.print_exc()

def load_terms():
    """Read the shared library fresh from disk so other users' additions show."""
    try:
        with open(TERMS_FILE, 'r', encoding='utf-8') as fh:
            d = json.load(fh)
            return d if isinstance(d, dict) else {}
    except (FileNotFoundError, ValueError, OSError):
        return {}

def add_terms(words, by, category=DEFAULT_CATEGORY):
    """Merge new approved technical terms into the shared file (atomic write).
    Each term is tagged with a category (item 5) for maintenance/auditing."""
    if category not in TERM_CATEGORIES:
        category = DEFAULT_CATEGORY
    with TERMS_LOCK:
        d = load_terms()
        added = []
        for w in words:
            wl = (w or '').strip().lower()
            if wl and wl not in d:
                d[wl] = {'by': by, 'at': datetime.now().strftime('%Y-%m-%d %H:%M'),
                         'category': category}
                added.append(wl)
        if added:
            _save_terms(d)
            for wl in added:
                log_term_event('add', wl, '', category)
        return d, added

def _save_terms(d):
    tmp = TERMS_FILE + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as fh:
        json.dump(d, fh, indent=2, ensure_ascii=False)
    os.replace(tmp, TERMS_FILE)

def move_terms(words, category):
    """Change the category of library words (audit trail kept in the history file)."""
    if category not in TERM_CATEGORIES:
        raise ValueError('Unknown category.')
    with TERMS_LOCK:
        d = load_terms()
        moved = []
        for w in words:
            wl = (w or '').strip().lower()
            if wl in d and (d[wl].get('category') or DEFAULT_CATEGORY) != category:
                frm = d[wl].get('category') or DEFAULT_CATEGORY
                d[wl]['category'] = category
                moved.append((wl, frm))
        if moved:
            _save_terms(d)
            for wl, frm in moved:
                log_term_event('move', wl, frm, category)
        return d, [m[0] for m in moved]

def remove_terms(words):
    with TERMS_LOCK:
        d = load_terms()
        removed = []
        for w in words:
            wl = (w or '').strip().lower()
            if wl in d:
                frm = d.pop(wl).get('category') or DEFAULT_CATEGORY
                removed.append((wl, frm))
        if removed:
            _save_terms(d)
            for wl, frm in removed:
                log_term_event('remove', wl, frm, '')
        return d, [r[0] for r in removed]

# ---------------------------------------------------------------- word logic
def deinflect(w):
    forms = {w}
    for suf, repl in (('ies','y'), ('ied','y'), ('ing',''), ('ed',''),
                      ('es',''), ('s',''), ('ly','')):
        if w.endswith(suf) and len(w) - len(suf) >= 2:
            base = w[:-len(suf)] + repl
            forms.add(base); forms.add(base + 'e')
    return forms

def lookup_not_approved(word):
    for f in deinflect(word):
        if f in NOT_APPROVED:
            return NOT_APPROVED[f]
    return None

def is_approved(word):
    return any(f in APPROVED for f in deinflect(word))

def is_self_reference(word, alts_list):
    """True when the dictionary's only 'alternative' is the same word
    (case-insensitive) - i.e. it is just marking the approved form,
    not asking for an actual replacement."""
    wl = word.strip().lower()
    return len(alts_list) == 1 and alts_list[0].strip().lower() == wl

# Months and ordinal suffixes cause false positives: "May" (month) is not the
# modal verb "may", and "6th" tokenizes to the fragment "th".
MONTHS = set('january february march april may june july august september '
             'october november december'.split())
ORDINAL_SUFFIX = {'st', 'nd', 'rd', 'th'}

def word_is_approved_form(word, alts_list):
    """True when the word already IS one of the approved alternatives
    (case-insensitive), so flagging it would just tell the author to replace a
    word with the same word - e.g. "related" when the alternative is RELATED,
    which is the correct STE past-participle form (Rule 3.3)."""
    wl = word.strip().lower()
    return any(wl == a.strip().lower() for a in alts_list)

def _pos_is_verb(pos):
    return bool(pos) and pos.strip().lower().split('/')[0].lstrip('.').startswith('v')

def alts_have_verb(alts_pos):
    return any(_pos_is_verb(p) for p in alts_pos)

# ---- sentence-context layer ---------------------------------------------------
_ROLE_MAP = {'n': 'n', 'tn': 'n', 'v': 'v', 'tv': 'v', 'adj': 'adj', 'adv': 'adv'}
_ROLE_NAME = {'n': 'noun', 'v': 'verb', 'adj': 'adjective', 'adv': 'adverb'}

def pos_roles(pos):
    """'n/v' -> {'n','v'}. Unknown or blank -> empty set."""
    out = set()
    for p in re.split(r'[\s/,]+', str(pos or '').lower()):
        p = p.strip('.()')
        if p in _ROLE_MAP:
            out.add(_ROLE_MAP[p])
    return out

def infer_role(prev, wi, mode):
    """Job of a word in the sentence, from the word before it. Returns
    'n', 'v', 'adj' or None when the sentence gives no clear sign."""
    cues = load_context_rules()['cues']
    if wi == 0:
        return 'v' if mode == 'procedural' else None      # a step starts with a verb
    if prev in cues.get('verb_before', ()):
        return 'v'
    if prev in cues.get('noun_before', ()):
        return 'n'
    if prev in cues.get('adj_before', ()):
        return 'adj'
    return None

def context_clear(wl, prev, nxt, wi, mode, na):
    """Decide from the sentence whether a flagged word is actually correct here.
    Returns a short reason text when the flag should be cleared, else None.
    Two checks: (1) author-added rules; (2) the word before shows the word is used
    in a role in which it IS approved (and not in the role that is not approved)."""
    ctx = load_context_rules()
    for r in ctx.get('rules', ()):
        if str(r.get('word', '')).strip().lower() != wl:
            continue
        pv = [str(x).lower() for x in (r.get('prev') or [])]
        nx = [str(x).lower() for x in (r.get('next') or [])]
        if not pv and not nx:
            continue
        if (not pv or prev in pv) and (not nx or nxt in nx):
            return 'Context rule %s%s' % (r.get('id') or '', (' - ' + r['note']) if r.get('note') else '')
    na_roles = pos_roles(na.get('pos', ''))
    ok_roles = pos_roles(approved_pos(wl))
    if not na_roles or not ok_roles:
        return None
    # "-ed" / "-ing" forms ("used", "mixed", "using") are verb forms, never nouns:
    # if the base verb is approved, the noun entry is not the one that applies.
    base = str(na.get('word', '')).lower()
    if wl != base and wl.endswith(('ed', 'ing')) and 'v' in ok_roles and 'v' not in na_roles:
        return '"%s" is a verb form of the approved verb "%s" (Rule 1.2)' % (wl, base)
    role = infer_role(prev, wi, mode)
    if role and role in ok_roles and role not in na_roles:
        return 'Used as a %s here; the word is approved as a %s (Rule 1.2)' % (
            _ROLE_NAME[role], ' / '.join(_ROLE_NAME[r] for r in sorted(ok_roles)))
    return None

# ---------------------------------------------------------------- rule engine
BRITISH = {
 'colour':'color','colours':'colors','behaviour':'behavior','behaviours':'behaviors',
 'centre':'center','centres':'centers','metre':'meter','metres':'meters','litre':'liter',
 'litres':'liters','fibre':'fiber','fibres':'fibers','defence':'defense','licence':'license',
 'catalogue':'catalog','catalogues':'catalogs','programme':'program','programmes':'programs',
 'analyse':'analyze','analysed':'analyzed','organise':'organize','organised':'organized',
 'realise':'realize','minimise':'minimize','maximise':'maximize','initialise':'initialize',
 'aluminium':'aluminum','tyre':'tire','tyres':'tires','grey':'gray','mould':'mold',
 'modelling':'modeling','labelling':'labeling','travelling':'traveling','cancelled':'canceled',
 'fuelled':'fueled','signalling':'signaling','practise':'practice','sulphur':'sulfur',
 'manoeuvre':'maneuver','oxidised':'oxidized','optimise':'optimize','favour':'favor',
}
CONTRACTIONS = {
 "don't":'do not',"doesn't":'does not',"didn't":'did not',"isn't":'is not',"aren't":'are not',
 "wasn't":'was not',"weren't":'were not',"can't":'cannot',"couldn't":'could not',
 "won't":'will not',"wouldn't":'would not',"shouldn't":'should not',"mustn't":'must not',
 "haven't":'have not',"hasn't":'has not',"hadn't":'had not',"it's":'it is',"that's":'that is',
 "there's":'there is',"we're":'we are',"you're":'you are',"they're":'they are',"i'm":'i am',
 "we've":'we have',"you've":'you have',"they've":'they have',"i've":'i have',"we'll":'we will',
 "you'll":'you will',"they'll":'they will',"it'll":'it will',"let's":'let us',"who's":'who is',
}
PHRASAL = {
 'set up':'install / prepare','carry out':'do','put out':'extinguish / remove','shut down':'stop',
 'pull out':'remove','take off':'remove','look for':'find','find out':'find','go through':'examine',
}
BE = {'is','are','was','were','be','been','being','am'}
HAVE = {'have','has','had'}
IRREG_PP = {'done','made','given','taken','seen','shown','set','put','cut','built','held','kept',
 'sent','found','removed','installed','completed','opened','closed','connected','disconnected',
 'gone','known','written','driven','broken','worn','torn','begun','shut','read','told'}

def _is_pp(w):
    return w.endswith('ed') or w in IRREG_PP

def analyze_constructs(s, mode):
    """Sentence-level rule findings that are NOT simple word swaps.
    Returns list of (rule, severity, message)."""
    out = []
    toks = WORD_RE.findall(s)
    low = [t.lower() for t in toks]
    if ';' in s:
        out.append(('8.1', 'error', 'Semicolon (;) is not allowed in STE.'))
    sl = ' ' + ' '.join(low) + ' '
    for ph, alt in PHRASAL.items():
        if ' ' + ph + ' ' in sl:
            out.append(('9.3', 'advisory', 'Possible phrasal verb "%s" - consider "%s".' % (ph, alt)))
    for i in range(len(low) - 1):
        a, b = low[i], low[i + 1]
        if a in BE and b.endswith('ing') and len(b) > 4:
            out.append(('3.5', 'warn', 'Progressive form "%s %s" - use the simple/imperative form.' % (a, b)))
        elif a in BE and _is_pp(b) and mode == 'procedural':
            out.append(('3.6', 'warn', 'Passive voice "%s %s" - use the active voice in procedures.' % (a, b)))
        if a in HAVE and _is_pp(b):
            out.append(('3.4', 'warn', 'Perfect tense "%s %s" - use the simple past tense.' % (a, b)))
    if mode == 'procedural' and low:
        if low[0] in ('you', 'we', 'the', 'this', 'these', 'it') and len(low) > 1 and \
           low[1] in ('must', 'should', 'will', 'can', 'shall', 'need', 'have', 'has'):
            out.append(('5.3', 'advisory', 'Write the instruction in the command (imperative) form.'))
    return out

# Full ASD-STE100 Issue 9 rule catalogue. mode: 'auto' = tool checks it,
# 'manual' = needs human judgement (listed so coverage is complete).
RULE_CATALOG = [
 ('Section 1 - Words', [
   ('1.1','Use approved words','auto'),
   ('1.2','Use words only as the specified part of speech','manual'),
   ('1.3','Use approved words only with their approved meanings','manual'),
   ('1.4','Use only approved forms of verbs and adjectives','auto'),
   ('1.5','Use words that fit a technical noun category','manual'),
   ('1.6','Use a non-approved word only as a technical noun','auto'),
   ('1.7','Do not use technical nouns as verbs','manual'),
   ('1.8','Use technical nouns approved in your company/industry','manual'),
   ('1.9','Select technical nouns that are short and clear','manual'),
   ('1.10','No regional, slang, or jargon technical nouns','manual'),
   ('1.11','Do not use different nouns for the same item','manual'),
   ('1.12','Use verbs that fit a technical verb category','manual'),
   ('1.13','Do not use technical verbs as nouns','manual'),
   ('1.14','Use American English spelling','auto'),
 ]),
 ('Section 2 - Multi-word nouns', [
   ('2.1','Multi-word nouns of no more than three words','manual'),
   ('2.2','Write nouns of more than three words in full','manual'),
 ]),
 ('Section 3 - Verbs', [
   ('3.1','Use only verb forms given in the dictionary','auto'),
   ('3.2','Use only the allowed verb forms and tenses','auto'),
   ('3.3','Use the past participle as an adjective','manual'),
   ('3.4','No auxiliary verbs for complex constructions','auto'),
   ('3.5','Use "-ing" only as a technical noun or modifier','auto'),
   ('3.6','Use the active voice (passive only in descriptive)','auto'),
   ('3.7','Use a verb, not a noun, to describe an action','manual'),
 ]),
 ('Section 4 - Sentences', [
   ('4.1','Write short and clear sentences','auto'),
   ('4.2','No omitted words or contractions','auto'),
   ('4.3','Use a vertical list for complex text','manual'),
   ('4.4','Use connecting words and phrases','manual'),
   ('4.5','Use an article or demonstrative adjective','manual'),
 ]),
 ('Section 5 - Procedural writing', [
   ('5.1','Maximum 20 words in a procedural sentence','auto'),
   ('5.2','Only one instruction in each sentence','manual'),
   ('5.3','Write instructions in the imperative form','auto'),
   ('5.4','Start with the condition, then the instruction','manual'),
   ('5.5','Write notes to give information, not instructions','manual'),
 ]),
 ('Section 6 - Descriptive writing', [
   ('6.1','Give information gradually','manual'),
   ('6.2','Use key words and phrases for structure','manual'),
   ('6.3','Maximum 25 words in a descriptive sentence','auto'),
   ('6.4','Use paragraphs to show related information','manual'),
   ('6.5','Each paragraph has only one topic','manual'),
   ('6.6','No paragraph has more than six sentences','auto'),
 ]),
 ('Section 7 - Safety instructions', [
   ('7.1','Identify with "warning" or "caution"','manual'),
   ('7.2','Start a safety instruction with a command','manual'),
   ('7.3','Give an explanation of the risk','manual'),
 ]),
 ('Section 8 - Punctuation and word count', [
   ('8.1','No semicolon; standard punctuation only','auto'),
   ('8.2','Use hyphens to connect related words','manual'),
   ('8.3','Use parentheses only as permitted','manual'),
   ('8.4','Colon in a vertical list counts like a period','manual'),
   ('8.5','Text in parentheses counts as one word','manual'),
   ('8.6','Count listed elements as one word','manual'),
   ('8.7','Hyphenated words count as one word','manual'),
 ]),
 ('Section 9 - Word usage', [
   ('9.1','Use a different construction when a swap fails','manual'),
   ('9.2','Use each approved word correctly','manual'),
   ('9.3','Do not make phrasal verbs','auto'),
   ('9.4','Use a consistent style and terminology','manual'),
 ]),
]

# ---------------------------------------------------------------- text parse
SENT_SPLIT = re.compile(r'(?<=[.!?:])\s+')
WORD_RE = re.compile(r"[A-Za-z][A-Za-z'\-]*")

def split_sentences(text):
    out = []
    para_i = 0
    for para in re.split(r'\n\s*\n', text):
        para = para.strip()
        if not para:
            continue
        sents = [s.strip() for s in SENT_SPLIT.split(para) if s.strip()]
        out.append((para_i, para, sents))
        para_i += 1
    return out

def default_replacement(alt):
    """Turn dictionary alternative into a natural body-text default."""
    if not alt or alt.startswith('Not in STE'):
        return ''
    return alt.lower()

def in_allowed(word, extra_allowed):
    return any(f in extra_allowed for f in deinflect(word))

def _para_snippet(para_text, word, radius=160):
    """Return the paragraph, or a window around the first match of `word`
    if the paragraph is long, so the review dropdown stays readable."""
    m = re.search(r"(?<![A-Za-z'])" + re.escape(word) + r"(?![A-Za-z'])",
                  para_text, re.IGNORECASE)
    if len(para_text) <= 2 * radius or not m:
        return para_text
    start = max(0, m.start() - radius)
    end = min(len(para_text), m.end() + radius)
    snippet = para_text[start:end]
    if start > 0:
        snippet = '...' + snippet
    if end < len(para_text):
        snippet = snippet + '...'
    return snippet

def check_text(text, mode, strict, extra_allowed=None):
    extra_allowed = extra_allowed or set()
    sent_limit = RULES['proc_sentence_max'] if mode == 'procedural' else RULES['desc_sentence_max']
    issues, flagged_words = [], {}
    rule_findings = []          # {rule, severity, msg, context}
    rule_counts = {}            # rule number -> count
    context_cleared = []        # flags cleared by the sentence-context layer
    total_words = total_sentences = 0
    paras = split_sentences(text)

    def add_finding(rule, sev, msg, context):
        rule_findings.append({'rule': rule, 'severity': sev, 'msg': msg, 'context': context})
        rule_counts[rule] = rule_counts.get(rule, 0) + 1

    paragraphs_text = {}

    for para_i, para_text, sents in paras:
        paragraphs_text[para_i] = para_text
        if len(sents) > RULES['para_max_sentences']:
            issues.append({'type': 'paragraph', 'severity': 'warn',
                           'msg': 'Paragraph has %d sentences (max %d).' %
                                  (len(sents), RULES['para_max_sentences']),
                           'context': (sents[0][:80] + '...') if sents else ''})
            add_finding('6.6', 'warn', 'Paragraph has %d sentences (max %d).' %
                        (len(sents), RULES['para_max_sentences']),
                        (sents[0][:80] + '...') if sents else '')
        for s in sents:
            total_sentences += 1
            ctx = s[:110] + ('...' if len(s) > 110 else '')
            words = WORD_RE.findall(s)
            total_words += len(words)
            if len(words) > sent_limit:
                issues.append({'type': 'length', 'severity': 'warn',
                               'msg': 'Sentence has %d words (max %d for %s writing).' %
                                      (len(words), sent_limit, mode), 'context': ctx})
                add_finding('5.1' if mode == 'procedural' else '6.3', 'warn',
                            'Sentence has %d words (max %d).' % (len(words), sent_limit), ctx)
            # sentence-level construct rules (verbs, semicolon, phrasal, imperative)
            for rule, sev, msg in analyze_constructs(s, mode):
                add_finding(rule, sev, msg, ctx)
            # word-level rules that ARE fixable word swaps
            for wi, w in enumerate(words):
                wl = w.lower()
                prev_w = words[wi - 1].lower() if wi > 0 else None
                next_w = words[wi + 1].lower() if wi + 1 < len(words) else None
                if wl in CONTRACTIONS:
                    e = flagged_words.setdefault('c:' + wl,
                        {'word': w, 'alt': CONTRACTIONS[wl], 'alts': [CONTRACTIONS[wl]],
                         'kind': 'rule', 'pos': '4.2', 'pos_display': 'Rule 4.2',
                         'rule': '4.2',
                         'guidance': 'Contractions are not allowed in STE - write "%s" in full '
                                     '(Rule 4.2).' % CONTRACTIONS[wl],
                         'count': 0, 'has_alt': True, 'replacement': CONTRACTIONS[wl], 'paras': set()})
                    e['count'] += 1
                    e['paras'].add(para_i)
                    rule_counts['4.2'] = rule_counts.get('4.2', 0) + 1
                    continue
                if wl in BRITISH and BRITISH[wl] != wl:
                    e = flagged_words.setdefault('b:' + wl,
                        {'word': w, 'alt': BRITISH[wl], 'alts': [BRITISH[wl]],
                         'kind': 'rule', 'pos': '1.14', 'pos_display': 'Rule 1.14',
                         'rule': '1.14',
                         'guidance': 'Use American English spelling "%s" (Rule 1.14).' % BRITISH[wl],
                         'count': 0, 'has_alt': True, 'replacement': BRITISH[wl], 'paras': set()})
                    e['count'] += 1
                    e['paras'].add(para_i)
                    rule_counts['1.14'] = rule_counts.get('1.14', 0) + 1
                    continue
                if wl in ALWAYS_OK:
                    continue
                # fix - ordinal fragments ("6th" -> "th") are not words
                if wl in ORDINAL_SUFFIX:
                    continue
                # fix - a capitalized month is a date, not the modal/verb sense
                # (e.g. "May 2026" is not the word "may")
                if wl in MONTHS and w[:1].isupper():
                    continue
                na = lookup_not_approved(wl)
                if na:
                    # context layer - clear the flag when the sentence shows the
                    # word is used in a role in which it is approved
                    why = context_clear(wl, prev_w, next_w, wi, mode, na)
                    if why:
                        if len(context_cleared) < 300:
                            context_cleared.append({'word': w, 'reason': why, 'context': ctx})
                        continue
                    # a word the team added to the shared library (or typed in the
                    # approved-terms box) is accepted, even if STE lists it as not approved
                    if in_allowed(wl, extra_allowed):
                        continue
                    alts_list = [a['alt'] for a in na.get('alts', []) if a.get('alt')]
                    alts_pos = [a.get('pos', '') for a in na.get('alts', []) if a.get('alt')]
                    if not alts_list and na.get('alt'):
                        alts_list = [na['alt']]
                    role_here = infer_role(prev_w, wi, mode)
                    wrong_role = bool(role_here) and role_here in pos_roles(na.get('pos', ''))
                    # a not-approved VERB used as a noun ("the oil") is a technical-noun
                    # question (Rule 1.5), not a wrong-verb error: offer the library.
                    if role_here == 'n' and pos_roles(na.get('pos', '')) == {'v'}:
                        if in_allowed(wl, extra_allowed) or not strict:
                            continue
                        e = flagged_words.setdefault('~' + wl,
                            {'word': w, 'alt': 'Used as a noun - technical noun?',
                             'alts': [], 'kind': 'unknown', 'pos': '', 'pos_display': '',
                             'rule': '1.5',
                             'guidance': ('"%s" is not approved as a verb, but here it is used as a '
                                          'noun. If it is a technical noun for your product, add it to '
                                          'the library (Rule 1.5). Otherwise use an approved noun.' % w.lower()),
                             'count': 0, 'has_alt': False, 'replacement': '', 'paras': set()})
                        e['count'] += 1
                        e['paras'].add(para_i)
                        rule_counts['1.5'] = rule_counts.get('1.5', 0) + 1
                        continue
                    # same word as its only alternative, e.g. test (v) -> TEST (n):
                    # it is a part-of-speech entry. Flag it only when the sentence
                    # shows the not-approved role.
                    if alts_list and is_self_reference(na['word'], alts_list) and not wrong_role:
                        continue
                    # fix - the word already IS the approved alternative
                    # (e.g. "related" when the alternative is RELATED); flagging
                    # it would suggest replacing a word with the same word.
                    # When the word is the base form itself ("test", "use") the
                    # sentence decides: flag only if it shows the not-approved role.
                    if alts_list and word_is_approved_form(w, alts_list):
                        if wl != na['word'].lower() or not wrong_role:
                            continue
                    # item 7 - part-of-speech aware: if the not-approved word is a
                    # verb but every alternative is a different part of speech
                    # (noun/adjective), a word-for-word swap breaks grammar
                    # ("investigate" -> "investigation"). Treat it as a reword,
                    # not a drop-in replacement, so we do not pre-fill a bad fix.
                    pos_mismatch = (_pos_is_verb(na.get('pos', '')) and alts_list
                                    and not alts_have_verb(alts_pos))
                    disp = ' / '.join(alts_list) if alts_list else '(rewrite - no direct alternative)'
                    # item 9 - detailed guidance with the approved alternative's
                    # usage and the governing rule reference.
                    if pos_mismatch:
                        rule_ref, has_fix, prefill = '9.1', False, ''
                        guidance = ('Not approved as a verb. The approved term is %s (%s) - '
                                    'reword the sentence rather than swapping the word '
                                    '(Rules 3.7 / 9.1).'
                                    % (disp, pos_label(alts_pos[0]) or 'different part of speech'))
                    elif alts_list:
                        rule_ref, has_fix = '1.1', True
                        prefill = default_replacement(alts_list[0])
                        guidance = ('Not approved in STE. Use %s instead (Rule 1.1 - use only '
                                    'approved words).' % disp)
                    else:
                        rule_ref, has_fix, prefill = '9.1', False, ''
                        guidance = ('Not approved in STE and has no direct replacement - '
                                    'reword the sentence (Rule 9.1).')
                    e = flagged_words.setdefault(na['word'].lower(),
                        {'word': na['word'], 'alt': disp, 'alts': alts_list, 'kind': 'dict',
                         'pos': na.get('pos', ''), 'pos_display': pos_label(na.get('pos', '')),
                         'alts_pos': alts_pos, 'rule': rule_ref, 'guidance': guidance,
                         'reword': pos_mismatch,
                         'count': 0, 'has_alt': has_fix,
                         'replacement': prefill, 'paras': set()})
                    e['count'] += 1
                    e['paras'].add(para_i)
                    rule_counts[rule_ref] = rule_counts.get(rule_ref, 0) + 1
                elif (strict and not is_approved(wl) and not w[0].isupper()
                      and not in_allowed(wl, extra_allowed)):
                    # item 9 - clearer guidance for words that are simply not in
                    # the dictionary (usually technical terms): tell the author
                    # exactly what to do and cite the governing rules.
                    guidance = ('Not found in the STE dictionary. If it is a technical noun or '
                                'verb for your product, add it to the technical library '
                                '(Rules 1.5 / 1.12). Otherwise replace it with an approved '
                                'word or reword (Rule 1.6).')
                    e = flagged_words.setdefault('~' + wl,
                        {'word': w, 'alt': 'Not in STE dictionary - approved technical term or reword?',
                         'alts': [], 'kind': 'unknown', 'pos': '', 'pos_display': '',
                         'rule': '1.6', 'guidance': guidance,
                         'count': 0, 'has_alt': False, 'replacement': '', 'paras': set()})
                    e['count'] += 1
                    e['paras'].add(para_i)
                    rule_counts['1.6'] = rule_counts.get('1.6', 0) + 1

    flagged = sorted(flagged_words.values(), key=lambda x: (-x['count'], x['word'].lower()))
    for f in flagged:
        para_ids = sorted(f.pop('paras', set()))
        f['occurrences'] = [{'para': i + 1, 'text': _para_snippet(paragraphs_text[i], f['word'])}
                             for i in para_ids]
    rule_findings.sort(key=lambda x: (x['rule']))

    # coverage report over the whole standard
    coverage = []
    auto_checked = auto_total = 0
    for section, rules in RULE_CATALOG:
        srules = []
        for num, title, kind in rules:
            cnt = rule_counts.get(num, 0)
            if kind == 'auto':
                auto_total += 1
                if cnt:
                    auto_checked += 1
            srules.append({'num': num, 'title': title, 'kind': kind, 'count': cnt})
        coverage.append({'section': section, 'rules': srules})

    summary = {
        'words': total_words, 'sentences': total_sentences, 'paragraphs': len(paras),
        'unapproved_total': sum(f['count'] for f in flagged),
        'length_issues': len([i for i in issues if i['type'] == 'length']),
        'paragraph_issues': len([i for i in issues if i['type'] == 'paragraph']),
        'rule_issues': len(rule_findings),
        'auto_rules': auto_total, 'mode': mode, 'sentence_limit': sent_limit,
    }
    return {'summary': summary, 'flagged': flagged, 'issues': issues,
            'rule_findings': rule_findings, 'coverage': coverage,
            'context_cleared': context_cleared,
            'highlighted': build_highlight(text, flagged_words, strict, extra_allowed, mode)}

def build_highlight(text, flagged_words, strict, extra_allowed=None, mode='procedural'):
    extra_allowed = extra_allowed or set()
    flagged_bases = {e['word'].lower() for e in flagged_words.values()}
    matches = list(WORD_RE.finditer(text))
    def neighbours(k):
        """(previous word, next word, position in sentence) for match k."""
        m = matches[k]
        prev = nxt = None
        wi = 0
        if k > 0 and not re.search(r'[.!?:\n]', text[matches[k - 1].end():m.start()]):
            prev = matches[k - 1].group(0).lower()
            wi = 1
        if k + 1 < len(matches) and not re.search(r'[.!?:\n]', text[m.end():matches[k + 1].start()]):
            nxt = matches[k + 1].group(0).lower()
        return prev, nxt, wi
    def repl(m, k):
        w = m.group(0); wl = w.lower()
        if wl in CONTRACTIONS:
            return '<mark class="bad" title="Write: %s">%s</mark>' % (
                html.escape(CONTRACTIONS[wl]), html.escape(w))
        if wl in BRITISH and BRITISH[wl] != wl:
            return '<mark class="bad" title="American: %s">%s</mark>' % (
                html.escape(BRITISH[wl]), html.escape(w))
        if in_allowed(wl, extra_allowed):
            return '<mark class="good" title="Approved technical term">%s</mark>' % html.escape(w)
        hit = lookup_not_approved(wl)
        if hit and hit['word'].lower() in flagged_bases:
            alts = [a['alt'] for a in hit.get('alts', []) if a.get('alt')] or ([hit['alt']] if hit.get('alt') else [])
            pv, nx, wi = neighbours(k)
            role = infer_role(pv, wi, mode)
            if alts and is_self_reference(hit['word'], alts) and not (
                    role and role in pos_roles(hit.get('pos', ''))):
                return html.escape(w)
            why = context_clear(wl, pv, nx, wi, mode, hit)
            if why:
                return '<span title="Cleared by context: %s">%s</span>' % (
                    html.escape(why, quote=True), html.escape(w))
            return '<mark class="bad" title="Use: %s">%s</mark>' % (
                html.escape(' / '.join(alts) or 'rewrite'), html.escape(w))
        if (strict and not is_approved(wl) and wl not in ALWAYS_OK
                and not w[0].isupper() and not in_allowed(wl, extra_allowed)):
            return '<mark class="warnw" title="Not in STE dictionary">%s</mark>' % html.escape(w)
        return html.escape(w)
    out, last = [], 0
    for k, m in enumerate(matches):
        out.append(html.escape(text[last:m.start()])); out.append(repl(m, k)); last = m.end()
    out.append(html.escape(text[last:]))
    return ''.join(out).replace('\n', '<br>')

# ---------------------------------------------------------------- readers
def _iter_row_cells(row):
    """Iterate a table row's actual cells straight from its XML, bypassing
    python-docx's row.cells property. That property computes each cell's
    position from the table's overall column grid and raises IndexError
    ('list index out of range') on many real-world documents whose tables
    have merged cells / an irregular grid - this avoids that entirely and,
    as a side effect, doesn't repeat a merged cell once per spanned column
    the way row.cells does."""
    from docx.table import _Cell
    for tc in row._tr.tc_lst:
        yield _Cell(tc, row.table)

def docx_units(doc):
    """Ordered list of 'replaceable units' (each a list of one or more
    docx Paragraph objects) matching 1:1 the paragraph numbering that
    check_text/split_sentences hand to the review UI: every top-level
    paragraph, then every table cell (a cell's paragraphs count as ONE
    unit, same as cell.text being one block). Empty ones are skipped.
    This exact order must be used everywhere a docx is opened so that
    unit index N always refers to the same location in the file."""
    units = []
    for p in doc.paragraphs:
        if p.text.strip():
            units.append([p])
    for t in doc.tables:
        for row in t.rows:
            for cell in _iter_row_cells(row):
                if cell.text.strip():
                    units.append(list(cell.paragraphs))
    return units

def xlsx_units(wb):
    """Ordered list of non-empty string cells, in the same worksheet/row/
    column order extract_text_from walks them for xlsx - so unit index N
    always means the same cell."""
    units = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    units.append(cell)
    return units

PARA_SPLIT = re.compile(r'(\n\s*\n)')

def txt_paragraph_blocks(text):
    """Split raw text/markdown into blocks, tagging each non-empty block
    with the paragraph index check_text/split_sentences would give it
    (blank-line separators and empty blocks are untagged)."""
    blocks, para_i = [], 0
    for i, part in enumerate(PARA_SPLIT.split(text)):
        if i % 2 == 1 or not part.strip():
            blocks.append({'text': part, 'para_i': None})
        else:
            blocks.append({'text': part, 'para_i': para_i})
            para_i += 1
    return blocks

# ------------------------------------------------------------------ item 4
# Page selection. "Pages" are defined per format:
#   .docx  -> content between rendered/explicit page breaks (Word's own pages)
#   .xlsx  -> each worksheet is one page
#   .txt   -> content between form-feed (\f) characters, else the whole file
# parse_page_range("1-3,5", 8) -> {1,2,3,5}
def parse_page_range(spec, n_pages):
    spec = (spec or '').strip()
    if not spec:
        return set(range(1, n_pages + 1))
    keep = set()
    for part in spec.split(','):
        part = part.strip()
        if not part:
            continue
        m = re.match(r'^(\d+)\s*-\s*(\d+)$', part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            for p in range(min(a, b), max(a, b) + 1):
                keep.add(p)
        elif part.isdigit():
            keep.add(int(part))
    return {p for p in keep if 1 <= p <= n_pages} or set(range(1, n_pages + 1))

def _docx_page_texts(path):
    """Return a list of page-text strings for a .docx, split on Word's page
    breaks (rendered or explicit). Always returns at least one page."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc = Document(path)
    pages = [[]]
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            p = Paragraph(child, doc)
            if p.text.strip():
                pages[-1].append(p.text)
            brs = child.findall('.//' + qn('w:br'))
            has_pb = any(b.get(qn('w:type')) == 'page' for b in brs) or \
                     bool(child.findall('.//' + qn('w:lastRenderedPageBreak')))
            if has_pb:
                pages.append([])
        elif child.tag == qn('w:tbl'):
            t = Table(child, doc)
            for row in t.rows:
                for cell in _iter_row_cells(row):
                    if cell.text.strip():
                        pages[-1].append(cell.text)
    texts = ['\n\n'.join(pg) for pg in pages]
    texts = [t for t in texts if t.strip()] or ['']
    return texts

def _xlsx_page_texts(path):
    """One page per worksheet."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages = []
    for ws in wb.worksheets:
        parts = []
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip():
                    parts.append(v)
        pages.append('\n\n'.join(parts))
    wb.close()
    return pages or ['']

def _txt_page_texts(path):
    with open(path, 'rb') as fh:
        raw = fh.read().decode('utf-8', errors='replace')
    return raw.split('\f') if '\f' in raw else [raw]

def page_texts(path, ext):
    if ext == '.docx':
        return _docx_page_texts(path)
    if ext in ('.xlsx', '.xlsm'):
        return _xlsx_page_texts(path)
    if ext in ('.txt', '.md'):
        return _txt_page_texts(path)
    return ['']

def page_count(path, ext):
    try:
        return len(page_texts(path, ext))
    except Exception:
        traceback.print_exc()
        return 1

def filter_text_pages(text, spec):
    """For pasted text, pages are form-feed separated (rare); usually 1 page."""
    pages = text.split('\f') if '\f' in text else [text]
    keep = parse_page_range(spec, len(pages))
    return '\n\n'.join(pages[i - 1] for i in sorted(keep) if 1 <= i <= len(pages))

def extract_text_from(path, ext, page_range=None):
    if page_range:
        pages = page_texts(path, ext)
        keep = parse_page_range(page_range, len(pages))
        return '\n\n'.join(pages[i - 1] for i in sorted(keep) if 1 <= i <= len(pages))
    if ext == '.docx':
        from docx import Document
        doc = Document(path)
        parts = ['\n'.join(pp.text for pp in unit) for unit in docx_units(doc)]
        return '\n\n'.join(parts)
    if ext in ('.txt', '.md'):
        with open(path, 'rb') as fh:
            return fh.read().decode('utf-8', errors='replace')
    if ext in ('.xlsx', '.xlsm'):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and v.strip():
                        parts.append(v)
        wb.close()
        return '\n\n'.join(parts)
    raise ValueError('Unsupported file type: ' + ext)

# ---------------------------------------------------------------- replacement
def cased(orig, repl):
    """Match the capitalization of the original occurrence."""
    if orig.isupper():
        return repl.upper()
    if orig[:1].isupper():
        return repl[:1].upper() + repl[1:]
    return repl

def build_pattern(find):
    return re.compile(r"(?<![A-Za-z'])" + re.escape(find) + r"(?![A-Za-z'])", re.IGNORECASE)

def pats_for_unit(reps, unit_i):
    """reps: list of {'find','replace','paras': set-of-int|None}. Returns the
    compiled (pattern, replacement) pairs that apply to paragraph/unit
    `unit_i` - i.e. reps with no paragraph restriction (apply everywhere)
    plus reps explicitly ticked for this unit."""
    return [(build_pattern(r['find']), r['replace']) for r in reps
            if r['find'] and r['replace'] and (r['paras'] is None or unit_i in r['paras'])]

def sub_text(text, pats, counter):
    for pat, rep in pats:
        def _r(m):
            counter[0] += 1
            return cased(m.group(0), rep)
        text = pat.sub(_r, text)
    return text

def replace_in_txt(src, dst, reps):
    counter = [0]
    with open(src, 'rb') as fh:
        text = fh.read().decode('utf-8', errors='replace')
    blocks = txt_paragraph_blocks(text)
    out = []
    for b in blocks:
        if b['para_i'] is None:
            out.append(b['text'])
            continue
        pats = pats_for_unit(reps, b['para_i'])
        out.append(sub_text(b['text'], pats, counter) if pats else b['text'])
    with open(dst, 'w', encoding='utf-8') as fh:
        fh.write(''.join(out))
    return counter[0]

def _replace_paragraph(p, pats, counter):
    """Apply replacements to a paragraph exactly once.
    Uses run-level replacement (keeps inline formatting) when every match sits
    inside a single run; falls back to a merged replacement only when a word is
    split across runs. This avoids double-applying when the replacement text
    itself contains the original word (e.g. motor -> dry-motor)."""
    runs = p.runs
    if not runs:
        return
    texts = [r.text for r in runs]
    joined = ''.join(texts)
    if not joined:
        return
    bounds, pos = [], 0
    for t in texts:
        bounds.append((pos, pos + len(t))); pos += len(t)
    cross = False
    for pat, _rep in pats:
        for mm in pat.finditer(joined):
            s, e = mm.span()
            if not any(bs <= s and e <= be for bs, be in bounds):
                cross = True; break
        if cross:
            break
    if cross:
        runs[0].text = sub_text(joined, pats, counter)
        for r in runs[1:]:
            r.text = ''
    else:
        for r in runs:
            if r.text:
                r.text = sub_text(r.text, pats, counter)

def replace_in_docx(src, dst, reps):
    from docx import Document
    counter = [0]
    doc = Document(src)
    for unit_i, paras in enumerate(docx_units(doc)):
        pats = pats_for_unit(reps, unit_i)
        if not pats:
            continue
        for p in paras:
            _replace_paragraph(p, pats, counter)
    # headers/footers aren't part of the paragraph numbering the review UI
    # shows, so only apply replacements that were NOT restricted to specific
    # paragraphs (i.e. "replace every occurrence") there.
    global_pats = [(build_pattern(r['find']), r['replace']) for r in reps
                   if r['find'] and r['replace'] and r['paras'] is None]
    if global_pats:
        try:
            for section in doc.sections:
                for hf in (section.header, section.footer):
                    for p in hf.paragraphs:
                        _replace_paragraph(p, global_pats, counter)
        except (AttributeError, Exception):
            pass
    doc.save(dst)
    return counter[0]

def replace_in_xlsx(src, dst, reps):
    import openpyxl
    counter = [0]
    wb = openpyxl.load_workbook(src)  # keep_vba false; formatting preserved
    for unit_i, cell in enumerate(xlsx_units(wb)):
        pats = pats_for_unit(reps, unit_i)
        if not pats:
            continue
        new = sub_text(cell.value, pats, counter)
        if new != cell.value:
            cell.value = new
    wb.save(dst)
    return counter[0]

# ---------------------------------------------------------------- annotation
def word_suggestion(w, strict, extra_allowed):
    """Suggestion string for one word, mirroring the on-screen check, or None."""
    wl = w.lower()
    if wl in CONTRACTIONS:
        return CONTRACTIONS[wl]
    if wl in BRITISH and BRITISH[wl] != wl:
        return BRITISH[wl]
    if wl in ALWAYS_OK or in_allowed(wl, extra_allowed):
        return None
    if wl in ORDINAL_SUFFIX:
        return None
    if wl in MONTHS and w[:1].isupper():
        return None
    na = lookup_not_approved(wl)
    if na:
        alts = [a['alt'] for a in na.get('alts', []) if a.get('alt')]
        alts_pos = [a.get('pos', '') for a in na.get('alts', []) if a.get('alt')]
        if not alts and na.get('alt'):
            alts = [na['alt']]
        if alts and is_self_reference(na['word'], alts):
            return None
        if alts and word_is_approved_form(w, alts):
            return None
        if _pos_is_verb(na.get('pos', '')) and alts and not alts_have_verb(alts_pos):
            return 'reword - approved term is %s' % ' / '.join(alts)
        return ' / '.join(alts) if alts else 'rewrite'
    if strict and not is_approved(wl) and not w[0].isupper():
        return 'verify: technical term or reword'
    return None

def _found_in_text(text, strict, extra):
    found, seen = [], set()
    for w in WORD_RE.findall(text):
        s = word_suggestion(w, strict, extra)
        if s and w.lower() not in seen:
            seen.add(w.lower())
            found.append((w, s))
    return found

def annotate_txt(src, dst, strict, extra):
    with open(src, 'rb') as fh:
        text = fh.read().decode('utf-8', errors='replace')
    n = [0]
    def repl(m):
        w = m.group(0)
        s = word_suggestion(w, strict, extra)
        if s:
            n[0] += 1
            return '%s [STE: %s]' % (w, s)
        return w
    with open(dst, 'w', encoding='utf-8') as fh:
        fh.write(WORD_RE.sub(repl, text))
    return n[0]

def _mark_run_inline(run, hits_map):
    """Fallback for python-docx builds without native comment support
    (< 1.1, e.g. the 0.8.6 seen in some locked-down environments): split
    `run` so ONLY the actually-flagged words get highlighted, each with its
    own short inline note right next to it - instead of highlighting the
    entire run and appending one giant list of every flagged word in it at
    the end (which is unreadable for long paragraphs)."""
    from docx.text.run import Run
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import RGBColor
    from copy import deepcopy

    text = run.text
    words_sorted = sorted(hits_map.keys(), key=len, reverse=True)
    pat = re.compile(r"(?<![A-Za-z'])(" + '|'.join(re.escape(w) for w in words_sorted) +
                      r")(?![A-Za-z'])", re.IGNORECASE)
    matches = list(pat.finditer(text))
    if not matches:
        return 0

    rPr_src = run._r.find(qn('w:rPr'))
    anchor = run._r
    parent = anchor.getparent()

    pieces, last = [], 0
    for m in matches:
        if m.start() > last:
            pieces.append(('text', text[last:m.start()], None))
        word = m.group(0)
        pieces.append(('flag', word, hits_map[word.lower()]))
        last = m.end()
    if last < len(text):
        pieces.append(('text', text[last:], None))

    XML_SPACE = qn('xml:space')

    def _text_run(value):
        r = OxmlElement('w:r')
        if rPr_src is not None:
            r.append(deepcopy(rPr_src))
        t = OxmlElement('w:t')
        t.set(XML_SPACE, 'preserve')
        t.text = value
        r.append(t)
        return r

    inserted = []
    for kind, value, sugg in pieces:
        r = _text_run(value)
        anchor.addprevious(r)
        inserted.append((kind, r, sugg))
    parent.remove(anchor)

    n = 0
    for kind, r, sugg in inserted:
        if kind != 'flag':
            continue
        Run(r, run.part).font.highlight_color = WD_COLOR_INDEX.YELLOW
        note_r = _text_run(' [STE: %s]' % sugg)
        r.addnext(note_r)
        note_run = Run(note_r, run.part)
        note_run.italic = True
        note_run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
        n += 1
    return n

def _annotate_paragraph(p, strict, extra, doc, mode):
    """Flag each occurrence of a not-approved word with a real Word comment
    anchored to just that run (mode[0] == 'native'), so the visible text is
    left exactly as-is and the reviewer sees/actions one comment per
    occurrence. Falls back to inline marking (see _mark_run_inline) when the
    installed python-docx is too old to support comments (< 1.1)."""
    n = 0
    for run in list(p.runs):
        if not run.text.strip():
            continue
        run_hits = _found_in_text(run.text, strict, extra)
        if not run_hits:
            continue
        if mode[0] == 'native':
            note = '; '.join('%s -> %s' % (w, s) for w, s in run_hits)
            try:
                doc.add_comment(runs=[run], text='STE: ' + note,
                                 author='STELOS', initials='STE')
                n += len(run_hits)
                continue
            except Exception:
                mode[0] = 'fallback'
        hits_map = {w.lower(): s for w, s in run_hits}
        n += _mark_run_inline(run, hits_map)
    return n

def _com_add_comments_for_terms(word_doc, terms):
    """terms: list of (find_text, comment_text). Mirrors the AuditSTE_Tool
    VBA macro's own pattern exactly: for each term, loop Word's native Find
    over the whole document and add a real comment at every match, then
    collapse to the end of that match and keep searching forward."""
    WD_FIND_STOP, WD_COLLAPSE_END = 0, 0
    n = 0
    for find_text, comment_text in terms:
        if not find_text:
            continue
        rng = word_doc.Content
        f = rng.Find
        f.ClearFormatting()
        f.Text = find_text
        f.MatchWholeWord = True
        f.MatchCase = False
        f.Forward = True
        f.Wrap = WD_FIND_STOP
        while f.Execute():
            try:
                word_doc.Comments.Add(Range=rng, Text=comment_text)
                n += 1
            except Exception:
                pass
            rng.Collapse(WD_COLLAPSE_END)
    return n

def annotate_docx_com(src, dst, strict, extra):
    """Add real Word comments via Word's own COM automation (win32com) -
    the exact same object model (doc.Comments.Add + Range.Find) as the
    working AuditSTE_Tool VBA macro. Used when the installed python-docx is
    too old for native add_comment support (< 1.1) but real Word is
    installed on the machine - true anywhere the VBA macro already runs."""
    import win32com.client as win32
    import pythoncom

    text = extract_text_from(src, '.docx')
    result = check_text(text, 'procedural', strict, extra)

    terms, seen = [], set()
    for f in result['flagged']:
        if f['word'].lower() in seen:
            continue
        seen.add(f['word'].lower())
        terms.append((f['word'], 'STE: %s -> %s' % (f['word'], f['alt'])))

    pythoncom.CoInitialize()
    word = None
    try:
        word = win32.DispatchEx('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0
        word_doc = word.Documents.Open(os.path.abspath(src), ReadOnly=False,
                                        AddToRecentFiles=False, ConfirmConversions=False)
        WD_FORMAT_XML_DOCUMENT = 12  # .docx
        try:
            n = _com_add_comments_for_terms(word_doc, terms)
            word_doc.SaveAs(os.path.abspath(dst), FileFormat=WD_FORMAT_XML_DOCUMENT)
        finally:
            word_doc.Close(SaveChanges=False)
        return n
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()

def annotate_docx(src, dst, strict, extra):
    from docx import Document
    doc = Document(src)

    if hasattr(doc, 'add_comment'):
        mode = ['native']
        n = 0
        for p in doc.paragraphs:
            n += _annotate_paragraph(p, strict, extra, doc, mode)
        for t in doc.tables:
            for row in t.rows:
                for cell in _iter_row_cells(row):
                    for p in cell.paragraphs:
                        n += _annotate_paragraph(p, strict, extra, doc, mode)
        doc.save(dst)
        return n, mode[0]

    # python-docx is too old for native comments - try driving real Word via
    # COM (win32com), same approach as the working AuditSTE_Tool VBA macro,
    # before falling back to inline highlight+note marking.
    try:
        n = annotate_docx_com(src, dst, strict, extra)
        return n, 'com'
    except Exception:
        pass

    mode = ['fallback']
    n = 0
    for p in doc.paragraphs:
        n += _annotate_paragraph(p, strict, extra, doc, mode)
    for t in doc.tables:
        for row in t.rows:
            for cell in _iter_row_cells(row):
                for p in cell.paragraphs:
                    n += _annotate_paragraph(p, strict, extra, doc, mode)
    doc.save(dst)
    return n, mode[0]

def annotate_xlsx(src, dst, strict, extra):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill
    fill = PatternFill('solid', fgColor='FFF3B0')
    wb = openpyxl.load_workbook(src)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value:
                    found = _found_in_text(cell.value, strict, extra)
                    if found:
                        note = 'STE: ' + '; '.join('%s -> %s' % (w, s) for w, s in found)
                        try:
                            cell.comment = Comment(note, 'STELOS')
                            cell.fill = fill
                        except Exception:
                            pass
                        n += len(found)
    wb.save(dst)
    return n
@app.route('/')
def index():
    return render_template_string(PAGE, meta=META, rules=RULES, max_upload_mb=MAX_UPLOAD_MB,
                                  ste_source=STE_SOURCE, ste_issue=STE_ISSUE,
                                  ste_issue_date=STE_ISSUE_DATE, ste_doc_url=STE_DOC_URL,
                                  term_categories=TERM_CATEGORIES,
                                  pilot_mode=PILOT_MODE, pilot_msg=PILOT_MSG)

@app.route('/check', methods=['POST'])
def check():
    try:
        mode = request.form.get('mode', 'procedural')
        strict = request.form.get('strict', 'true') == 'true'
        raw_terms = request.form.get('terms', '') or ''
        extra_allowed = set()
        for t in re.split(r'[\n,;]+', raw_terms):
            t = t.strip().lower()
            if t:
                extra_allowed.add(t)
        library = load_terms()                 # shared technical-terms DB
        extra_allowed |= set(library.keys())
        text = request.form.get('text', '') or ''
        token = None
        f = request.files.get('file')
        if f and f.filename:
            ext = os.path.splitext(f.filename)[1].lower()
            if ext not in ('.docx', '.txt', '.md', '.xlsx', '.xlsm'):
                return jsonify({'error': 'Unsupported file. Use .docx, .xlsx, .txt or .md, or paste text.'}), 400
            token = uuid.uuid4().hex
            path = os.path.join(TMPDIR, token + ext)
            f.save(path)
            UPLOADS[token] = {'path': path, 'filename': f.filename, 'ext': ext}
            text = extract_text_from(path, ext)
        elif text.strip():
            # store pasted text so replacement can produce a corrected copy too
            token = uuid.uuid4().hex
            path = os.path.join(TMPDIR, token + '.txt')
            with open(path, 'w', encoding='utf-8') as fh:
                fh.write(text)
            UPLOADS[token] = {'path': path, 'filename': 'pasted_text.txt', 'ext': '.txt'}
        if not text.strip():
            return jsonify({'error': 'No text provided. Paste text or upload a file.'}), 400
        # item 4 - restrict to selected pages if the user asked
        page_range = (request.form.get('pages', '') or '').strip()
        if page_range and token in UPLOADS:
            text = extract_text_from(UPLOADS[token]['path'], UPLOADS[token]['ext'],
                                     page_range) or text
        elif page_range:
            text = filter_text_pages(text, page_range)
        # item 1 - repair split words / invisible characters before checking
        wb_notes = []
        text = normalize_text(text, wb_notes)
        result = check_text(text, mode, strict, extra_allowed)
        result['word_break_notes'] = wb_notes
        result['pilot'] = PILOT_MODE
        threading.Thread(target=update_corpus, args=(WORD_RE.findall(text),), daemon=True).start()
        result['token'] = token
        result['library'] = library
        result['user'] = CURRENT_USER
        doc_name = UPLOADS[token]['filename'] if token in UPLOADS else 'pasted text'
        s = result['summary']
        log_usage('check', documents=doc_name, mode=mode, words=s['words'],
                   sentences=s['sentences'], unapproved=s['unapproved_total'],
                   rule_findings=s['rule_issues'])
        return jsonify(result)
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/pages', methods=['POST'])
def pages():
    """Report how many 'pages' an uploaded file has, so the UI can offer page
    selection (item 4). Called when the user picks a file."""
    try:
        f = request.files.get('file')
        if not f or not f.filename:
            return jsonify({'error': 'No file.'}), 400
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ('.docx', '.txt', '.md', '.xlsx', '.xlsm'):
            return jsonify({'error': 'Unsupported file.'}), 400
        tmp = os.path.join(TMPDIR, uuid.uuid4().hex + ext)
        f.save(tmp)
        try:
            pts = page_texts(tmp, ext)
            unit = 'sheet' if ext in ('.xlsx', '.xlsm') else 'page'
            labels = []
            for i, t in enumerate(pts, 1):
                words = len(WORD_RE.findall(t))
                labels.append({'page': i, 'words': words})
        finally:
            try:
                os.remove(tmp)
            except OSError:
                pass
        return jsonify({'count': len(pts), 'unit': unit, 'labels': labels})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/apply', methods=['POST'])
def apply():
    try:
        if PILOT_MODE:
            return jsonify({'error': PILOT_MSG}), 403
        data = request.get_json(force=True)
        token = data.get('token')
        reps = []
        for p in data.get('replacements', []):
            find, rep = p.get('find', ''), p.get('replace', '')
            if not find or not rep:
                continue
            paras = p.get('paras')  # list of paragraph indices to restrict to, or falsy = everywhere
            reps.append({'find': find, 'replace': rep,
                         'paras': set(paras) if paras else None})
        if token not in UPLOADS:
            return jsonify({'error': 'Session expired - run the check again.'}), 400
        if not reps:
            return jsonify({'error': 'No replacements selected (tick words and set a replacement).'}), 400

        info = UPLOADS[token]
        src, ext = info['path'], info['ext']
        stem = os.path.splitext(os.path.basename(info['filename']))[0]
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        out_name = '%s_STE_%s%s' % (stem, ts, ext)
        dst = os.path.join(TMPDIR, out_name)

        if ext == '.docx':
            n = replace_in_docx(src, dst, reps)
        elif ext in ('.xlsx', '.xlsm'):
            n = replace_in_xlsx(src, dst, reps)
        else:
            n = replace_in_txt(src, dst, reps)

        try:
            resp = send_file(dst, as_attachment=True, download_name=out_name)      # Flask 2.x
        except TypeError:
            resp = send_file(dst, as_attachment=True, attachment_filename=out_name)  # Flask 1.1.x
        resp.headers['X-STE-Replacements'] = str(n)
        resp.headers['X-STE-Outfile'] = out_name
        resp.headers['Access-Control-Expose-Headers'] = 'X-STE-Replacements, X-STE-Outfile'
        log_usage('download', documents=info['filename'], replacements=n)
        return resp
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/annotate', methods=['POST'])
def annotate():
    """Produce an annotated copy of the ORIGINAL file (same format) with STE
    suggestions marked in-line. The original is never changed."""
    try:
        if PILOT_MODE:
            return jsonify({'error': PILOT_MSG}), 403
        data = request.get_json(force=True)
        token = data.get('token')
        strict = bool(data.get('strict', True))
        raw_terms = data.get('terms', '') or ''
        if token not in UPLOADS:
            return jsonify({'error': 'Session expired - run the check again.'}), 400
        extra = set()
        for t in re.split(r'[\n,;]+', raw_terms):
            t = t.strip().lower()
            if t:
                extra.add(t)
        extra |= set(load_terms().keys())

        info = UPLOADS[token]
        src, ext = info['path'], info['ext']
        stem = os.path.splitext(os.path.basename(info['filename']))[0]
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        out_name = '%s_STE_annotated_%s%s' % (stem, ts, ext)
        dst = os.path.join(TMPDIR, out_name)

        mode = 'native'
        if ext == '.docx':
            n, mode = annotate_docx(src, dst, strict, extra)
        elif ext in ('.xlsx', '.xlsm'):
            n = annotate_xlsx(src, dst, strict, extra)
        else:
            n = annotate_txt(src, dst, strict, extra)

        try:
            resp = send_file(dst, as_attachment=True, download_name=out_name)
        except TypeError:
            resp = send_file(dst, as_attachment=True, attachment_filename=out_name)
        resp.headers['X-STE-Annotations'] = str(n)
        resp.headers['X-STE-Outfile'] = out_name
        resp.headers['X-STE-Comment-Mode'] = mode
        resp.headers['Access-Control-Expose-Headers'] = \
            'X-STE-Annotations, X-STE-Outfile, X-STE-Comment-Mode'
        log_usage('annotate', documents=info['filename'], annotations=n)
        return resp
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/terms', methods=['GET'])
def terms():
    return jsonify({'library': load_terms(), 'user': CURRENT_USER,
                    'categories': TERM_CATEGORIES})

@app.route('/add_term', methods=['POST'])
def add_term():
    try:
        data = request.get_json(force=True)
        words = data.get('terms', [])
        category = data.get('category', DEFAULT_CATEGORY)
        if not words:
            return jsonify({'error': 'No words selected.'}), 400
        library, added = add_terms(words, CURRENT_USER, category)
        return jsonify({'ok': True, 'added': added, 'library': library,
                        'user': CURRENT_USER, 'categories': TERM_CATEGORIES})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/move_term', methods=['POST'])
def move_term():
    try:
        data = request.get_json(force=True)
        words = data.get('terms', [])
        if not words:
            return jsonify({'error': 'Tick at least one word.'}), 400
        library, moved = move_terms(words, data.get('category', ''))
        return jsonify({'ok': True, 'moved': moved, 'library': library, 'user': CURRENT_USER})
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/remove_term', methods=['POST'])
def remove_term():
    try:
        data = request.get_json(force=True)
        words = data.get('terms', [])
        if not words:
            return jsonify({'error': 'Tick at least one word.'}), 400
        library, removed = remove_terms(words)
        return jsonify({'ok': True, 'removed': removed, 'library': library, 'user': CURRENT_USER})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/terms_export', methods=['GET'])
def terms_export():
    """Whole shared library as a CSV (Word, Category, Added by, Added at) for audit."""
    import io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(['Word', 'Category', 'Added by', 'Added at'])
    for k, v in sorted(load_terms().items()):
        w.writerow([k, (v or {}).get('category') or DEFAULT_CATEGORY,
                    (v or {}).get('by', ''), (v or {}).get('at', '')])
    from flask import Response
    return Response('﻿' + buf.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename=stelos_library_%s.csv'
                             % datetime.now().strftime('%Y%m%d')})

@app.route('/terms_history', methods=['GET'])
def terms_history():
    return jsonify({'history': load_term_history(200)})

@app.route('/pos_ref', methods=['GET'])
def pos_ref():
    return jsonify(load_pos_ref())

@app.route('/harvest', methods=['POST'])
def harvest():
    """Library-building: read one or more documents and list the words that are
    in neither the STE dictionary nor the shared library, with an example
    sentence, so Authors / CAs can sort them into categories."""
    try:
        files = [f for f in request.files.getlist('files') if f and f.filename]
        pasted = request.form.get('text', '') or ''
        page_range = (request.form.get('pages', '') or '').strip()
        if not files and not pasted.strip():
            return jsonify({'error': 'Choose one or more documents, or paste text.'}), 400
        library = load_terms()
        counts, example, caps, docs = {}, {}, {}, []
        all_words = []
        sources = [(f.filename, f) for f in files]
        if pasted.strip():
            sources.append(('pasted text', None))
        for name, f in sources:
            if f is None:
                text = pasted
            else:
                ext = os.path.splitext(name)[1].lower()
                if ext not in ('.docx', '.txt', '.md', '.xlsx', '.xlsm'):
                    return jsonify({'error': 'Unsupported file: %s. Use .docx, .xlsx, .txt or .md.' % name}), 400
                path = os.path.join(TMPDIR, uuid.uuid4().hex + ext)
                f.save(path)
                try:
                    text = extract_text_from(path, ext, page_range or None)
                finally:
                    try:
                        os.remove(path)
                    except OSError:
                        pass
            text = normalize_text(text)
            docs.append(name)
            all_words.extend(WORD_RE.findall(text))
            for _, _, sents in split_sentences(text):
                for s in sents:
                    for w in WORD_RE.findall(s):
                        wl = w.lower().strip("'-")
                        if (len(wl) < 3 or not wl.replace('-', '').isalpha() or wl in ALWAYS_OK
                                or wl in CONTRACTIONS or wl in MONTHS or wl in library
                                or is_approved(wl) or lookup_not_approved(wl)
                                or in_allowed(wl, set(library))):
                            continue
                        counts[wl] = counts.get(wl, 0) + 1
                        if w[:1].isupper():
                            caps[wl] = caps.get(wl, 0) + 1
                        if wl not in example:
                            example[wl] = s[:160]
        update_corpus(all_words)
        rows = [{'word': w, 'count': c, 'capitalised': caps.get(w, 0) > 0, 'example': example[w]}
                for w, c in sorted(counts.items(), key=lambda x: (-x[1], x[0]))][:1000]
        log_usage('harvest', documents=', '.join(docs)[:300], doc_count=len(docs),
                  words=len(all_words), unapproved=len(counts))
        return jsonify({'rows': rows, 'total_unknown': len(counts), 'documents': docs,
                        'words': len(all_words)})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

# ---------------------------------------------------------------- template
PAGE = r"""<!doctype html><html><head><meta charset="utf-8">
<title>STELOS - The ASD-STE100 Intelligence Platform</title>
<link rel="icon" href="https://www.alten.com/wp-content/uploads/2019/01/favicon-alten.png">
<style>
*,*::before,*::after{box-sizing:border-box;margin:0;padding:0}
:root{--bg:#eef2f9;--surface:#fff;--surface2:#f1f5fb;--border:#c8d4e8;--accent:#1a4fad;
--accent2:#1d6fdb;--green:#059669;--red:#dc2626;--amber:#d97706;--text:#1e293b;--muted:#475569;
--mono:'IBM Plex Mono',Consolas,monospace}
body{font-family:system-ui,'Segoe UI',sans-serif;font-size:14px;background:var(--bg);color:var(--text);
min-height:100vh;display:flex;flex-direction:column}
header{background:var(--accent);color:#fff;padding:14px 28px;display:flex;align-items:center;gap:20px;
box-shadow:0 2px 8px rgba(0,0,0,.18)}
.header-logo{display:flex;align-items:center;gap:14px}
.header-logo img{height:34px;width:auto;background:#fff;border-radius:6px;padding:4px}
.brand-name{font-size:24px;font-weight:800;letter-spacing:1.5px}
.header-title h1{font-size:20px;font-weight:600}.header-title .subtitle{font-size:12px;opacity:.85;margin-top:2px}
.brand-sep{width:1px;height:34px;background:rgba(255,255,255,.35)}
.brand-by{font-size:11px;opacity:.85;margin-top:3px;letter-spacing:.5px}
main{flex:1;max-width:1040px;width:100%;margin:22px auto;padding:0 20px}
.card{background:var(--surface);border:1px solid var(--border);border-radius:10px;padding:18px 20px;margin-bottom:18px;
box-shadow:0 1px 3px rgba(0,0,0,.05)}
.card h2{font-size:16px;margin-bottom:12px;color:var(--accent)}
textarea{width:100%;min-height:170px;border:1px solid var(--border);border-radius:8px;padding:12px;
font-family:var(--mono);font-size:13px;resize:vertical}
.row{display:flex;gap:16px;flex-wrap:wrap;align-items:center;margin-top:12px}
label.opt{display:flex;align-items:center;gap:6px;cursor:pointer}
select,input[type=file],input.rep{padding:7px 9px;border:1px solid var(--border);border-radius:6px;background:#fff;font-size:13px}
input.rep{font-family:var(--mono);width:100%;color:var(--green)}
select.repsel{font-family:var(--mono);font-size:12px;padding:5px 6px;border:1px solid var(--border);
border-radius:6px;background:#f0fdf4;color:#166534;margin-bottom:5px;max-width:100%}
button{background:var(--accent2);color:#fff;border:0;padding:10px 20px;border-radius:7px;font-size:14px;
font-weight:600;cursor:pointer}button:disabled{opacity:.5;cursor:not-allowed}
button.sec{background:var(--surface2);color:var(--accent);border:1px solid var(--border)}
button.go{background:var(--green)}
.stats{display:flex;gap:14px;flex-wrap:wrap}
.stat{background:var(--surface2);border:1px solid var(--border);border-radius:8px;padding:10px 14px;min-width:110px}
.stat .n{font-size:22px;font-weight:700}.stat .l{font-size:11px;color:var(--muted);text-transform:uppercase;letter-spacing:.4px}
.stat.bad .n{color:var(--red)}.stat.warn .n{color:var(--amber)}.stat.ok .n{color:var(--green)}
table{width:100%;border-collapse:collapse;margin-top:6px}
th,td{text-align:left;padding:8px 10px;border-bottom:1px solid var(--border);font-size:13px;vertical-align:middle}
th{background:var(--surface2);color:var(--muted);font-size:11px;text-transform:uppercase;letter-spacing:.4px}
td.w{font-family:var(--mono);color:var(--red);font-weight:600}
td.chk{width:34px;text-align:center}td.repcell{width:230px}
mark.bad{background:#fde2e2;color:#991b1b;border-radius:3px;padding:0 2px;cursor:help;border-bottom:2px solid var(--red)}
mark.warnw{background:#fef3c7;color:#92400e;border-radius:3px;padding:0 2px;cursor:help}
.doc{background:var(--surface2);border:1px solid var(--border);border-radius:8px;padding:14px;line-height:1.7;
font-size:13.5px;max-height:340px;overflow:auto}
.issue{padding:8px 12px;border-left:3px solid var(--amber);background:#fffbeb;border-radius:4px;margin-bottom:8px;font-size:13px}
.issue .ctx{color:var(--muted);font-family:var(--mono);font-size:12px;margin-top:3px}
.issue.sev-error{border-left-color:var(--red);background:#fef2f2}
.issue.sev-advisory{border-left-color:var(--accent2);background:#eff6ff}
.rulebadge{display:inline-block;background:var(--accent);color:#fff;font-size:11px;font-weight:700;
border-radius:4px;padding:1px 7px;margin-right:6px}
.sevtag{font-size:10px;font-weight:700;letter-spacing:.4px;padding:1px 6px;border-radius:3px;margin-right:6px}
.sevtag.error{background:#fde2e2;color:#991b1b}.sevtag.warn{background:#fef3c7;color:#92400e}
.sevtag.advisory{background:#dbeafe;color:#1e40af}
.covsec{margin-bottom:14px}.covsec h4{font-size:13px;color:var(--muted);margin-bottom:6px;
text-transform:uppercase;letter-spacing:.4px}
.covgrid{display:grid;grid-template-columns:1fr 1fr;gap:6px}
.covitem{font-size:12.5px;padding:6px 9px;border-radius:6px;border:1px solid var(--border);
display:flex;align-items:center;gap:6px}
.covitem b{font-family:var(--mono)}
.covitem .covtag{margin-left:auto;font-size:10px;font-weight:700;padding:1px 6px;border-radius:10px}
.cov-ok{background:#f0fdf4;border-color:#bbf7d0}.cov-ok .covtag{background:#bbf7d0;color:#166534}
.cov-hit{background:#fef2f2;border-color:#fecaca}.cov-hit .covtag{background:#fecaca;color:#991b1b}
.cov-man{background:var(--surface2)}.cov-man .covtag{background:#e2e8f0;color:#475569}
@media(max-width:640px){.covgrid{grid-template-columns:1fr}}
.hidden{display:none}.err{color:var(--red);font-weight:600;margin-top:10px}
.ok-msg{color:var(--green);font-weight:600;margin-top:10px}
footer{text-align:center;padding:16px;font-size:12px;color:var(--muted)}
.tag{display:inline-block;font-size:11px;background:var(--surface2);border:1px solid var(--border);
border-radius:12px;padding:2px 10px;color:var(--muted);margin-left:6px}
.note{font-size:12px;color:var(--muted);margin:4px 0 10px}
details.guide{background:#eff6ff;border:1px solid #bfdbfe;border-radius:8px;padding:8px 12px;margin-bottom:12px}
details.guide>summary{cursor:pointer;color:var(--accent);font-weight:600;font-size:13px}
.guidebox{margin-top:8px;font-size:13px;line-height:1.55}.guidebox p{margin-bottom:7px}
.examples{display:flex;flex-direction:column;gap:6px;margin:6px 0}
.ex{font-size:12.5px;padding:6px 9px;border-radius:6px;font-family:var(--mono)}
.ex.ok{background:#f0fdf4;border:1px solid #bbf7d0}.ex.bad{background:#fef2f2;border:1px solid #fecaca}
.extag{display:inline-block;font-size:10px;font-weight:700;padding:1px 6px;border-radius:10px;margin-right:6px;font-family:system-ui}
.extag.good{background:#bbf7d0;color:#166534}.extag.bad{background:#fecaca;color:#991b1b}
.guidance{font-size:11.5px;color:var(--muted);margin-top:3px;line-height:1.45}
.rewordtag{display:inline-block;background:#fef3c7;color:#92400e;font-size:11px;font-weight:700;
border-radius:6px;padding:3px 9px}
.posbadge{display:inline-block;background:#e0e7ff;color:#3730a3;font-size:10px;font-weight:700;
border-radius:4px;padding:1px 6px;font-family:system-ui}
.libcat{margin-bottom:10px}.libcat h5{font-size:12px;color:var(--muted);text-transform:uppercase;
letter-spacing:.4px;margin-bottom:5px}
mark.good{background:#dcfce7;color:#166534;border-radius:3px;padding:0 2px;border-bottom:2px solid var(--green)}
button.tech{background:var(--green)}
.techcell label{display:flex;align-items:center;gap:6px;font-size:12px;color:var(--muted);cursor:pointer}
.lib{margin-top:14px;border-top:1px dashed var(--border);padding-top:12px}
.lib h4{font-size:13px;color:var(--accent);margin-bottom:8px}
.chips{display:flex;flex-wrap:wrap;gap:6px}
.chip{background:#dcfce7;color:#166534;border:1px solid #bbf7d0;border-radius:14px;padding:2px 10px;
font-size:12px;font-family:var(--mono)}.chip small{color:#3f7d55;font-family:system-ui}
.rowdone{opacity:.55}
.paratoggle{background:none;border:1px solid var(--border);border-radius:4px;color:var(--accent);
font-size:11px;padding:1px 5px;cursor:pointer;margin-right:4px}
.paralistrow td{border-bottom:1px solid var(--border);background:var(--surface2);padding:8px 10px 10px 10px}
.paralist{display:flex;flex-direction:column;gap:6px}
.paraitem{font-size:12.5px;line-height:1.6;background:var(--surface);border:1px solid var(--border);
border-radius:6px;padding:6px 10px;display:flex;align-items:flex-start;gap:8px}
.paraitem b{color:var(--muted);font-weight:600;margin-right:4px}
.paraitemchk{display:flex;align-items:center;padding-top:2px;cursor:pointer}
.paraitem.unchecked{opacity:.45}
.pilot{background:#fff4d6;border-bottom:2px solid #e0a800;color:#5a4300;padding:9px 24px;font-size:13px}
.pilot b{color:#3d2e00}
.wbnote{background:#fff4d6;border:1px solid #e0a800;border-radius:8px;padding:10px 14px;margin-bottom:14px;font-size:13px;color:#5a4300}
.ctxrow{font-size:12.5px;padding:6px 0;border-bottom:1px dashed var(--border)}
.ctxrow .why{color:var(--green);font-size:12px}
.ctxrow .snip{color:var(--muted);font-size:12px}
.audbar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:8px 0}
.audbar select,.audbar input[type=text]{padding:6px 8px;border:1px solid var(--border);border-radius:6px;font-size:12px}
.chip.sel{outline:2px solid var(--accent)}
.chip input{margin-right:4px;vertical-align:middle}
.postbl{width:100%;border-collapse:collapse;font-size:13px}
.postbl th,.postbl td{border:1px solid var(--border);padding:7px 9px;text-align:left;vertical-align:top}
.postbl th{background:var(--surface)}
.postbl .abbr{font-family:var(--mono);font-weight:700;white-space:nowrap}
.postbl .stx{color:var(--green)} .postbl .nstx{color:#b3261e}
.hrow td{font-size:12px}
a.qlink{display:inline-block;margin-left:4px;width:16px;height:16px;line-height:16px;text-align:center;border-radius:50%;
  background:var(--accent);color:#fff;font-size:11px;text-decoration:none;font-weight:700}
</style></head><body>
<header>
  <div class="header-logo">
    <img src="https://www.alten.com/wp-content/uploads/2019/01/favicon-alten.png" alt="ALTEN"
         onerror="this.replaceWith(Object.assign(document.createElement('span'),{textContent:'ALTEN',style:'font-weight:800;background:#fff;color:#1a4fad;padding:6px 10px;border-radius:6px'}))">
    <span class="brand-sep"></span>
    <span class="brand-name">STELOS</span>
  </div>
  <div class="header-title"><h1>The ASD-STE100 Intelligence Platform</h1>
  <p class="brand-by">An ALTEN product &middot; Validating against {{ste_source}}
    {% if ste_doc_url %}&middot; <a href="{{ste_doc_url}}" target="_blank"
       style="color:#fff;text-decoration:underline">specification</a>{% endif %}</p></div>
</header>
{% if pilot_mode %}<div class="pilot"><b>PILOT &mdash; library building.</b> {{pilot_msg}}
  Upload documents in <a href="#libBuilder">Library builder</a> to feed new words into the shared library.</div>{% endif %}
<main>
  <div class="card">
    <h2>1 &middot; Input <span class="tag">{{meta.approved_count}} approved &middot; {{meta.not_approved_count}} not approved</span>
      <span class="tag" title="{{ste_source}}">Validation: {{ste_source}}</span></h2>

    <details class="guide" open>
      <summary>&#9432; What to enter &amp; examples (click to hide)</summary>
      <div class="guidebox">
        <p><b>What this tool checks:</b> your finished technical writing (a procedure, a description,
        a caption, a note) against ASD-STE100 Simplified Technical English &mdash; approved words,
        sentence length, verb forms, punctuation and spelling.</p>
        <p><b>What to enter:</b> paste the actual sentences you wrote, or upload the document.
        Enter real, complete sentences &mdash; not single words, code, tables of part numbers, or
        file names.</p>
        <div class="examples">
          <div class="ex ok"><span class="extag good">Good input</span>
            Remove the bolt. Install the new seal. Do a test of the system for leaks.</div>
          <div class="ex bad"><span class="extag bad">Avoid</span>
            single words (&ldquo;valve&rdquo;), part-number lists, or raw code &mdash; these are not sentences
            and will over-flag.</div>
        </div>
        <p class="note" style="margin-bottom:0">Pick <b>Procedural</b> for instructions (steps, commands)
        and <b>Descriptive</b> for explanations. Add your product's technical terms below so they are
        not flagged.</p>
      </div>
    </details>

    <textarea id="text" placeholder="Paste real, complete sentences here. Example: Remove the bolt. Install the new seal."></textarea>
    <div class="row">
      <input type="file" id="file" accept=".docx,.xlsx,.xlsm,.txt,.md" onchange="onFilePick()">
      <label class="opt">Writing type:
        <select id="mode">
          <option value="procedural">Procedural (max {{rules.proc_sentence_max}} words/sentence)</option>
          <option value="descriptive">Descriptive (max {{rules.desc_sentence_max}} words/sentence)</option>
        </select>
      </label>
      <label class="opt"><input type="checkbox" id="strict" checked>
        Full ASD-STE100 check (flag <b>every</b> word not in the approved dictionary)</label>
    </div>

    <div class="row" id="pageRow" style="display:none">
      <span id="pageInfo" class="note" style="margin:0"></span>
      <label class="opt">Pages to check
        <input type="text" id="pages" placeholder="all (e.g. 1-3,5)" style="width:130px;
          padding:7px 9px;border:1px solid var(--border);border-radius:6px;font-size:13px">
      </label>
    </div>

    <p class="note">Accepted files: Word (.docx), Excel (.xlsx / .xlsm), text (.txt / .md), up to {{max_upload_mb}} MB.
    The original file is never changed &mdash; replacements are written to a downloaded copy.
    STE is a closed vocabulary: with the full check on, any word outside the approved list is
    flagged as <i>"Not in STE dictionary - verify it is an approved technical noun/verb or reword."</i>
    Turn it off to see only words that have a suggested replacement.</p>
    <details style="margin-top:6px">
      <summary style="cursor:pointer;color:var(--accent);font-size:13px">Approved technical terms (optional)</summary>
      <p class="note">Add your project's legitimate technical nouns/verbs (valve, sensor, zone, part numbers...),
      one per line or comma-separated. These will not be flagged by the full check.</p>
      <textarea id="terms" style="min-height:70px" placeholder="valve&#10;sensor&#10;zone&#10;solenoid"></textarea>
    </details>
    <div class="row">
      <button id="runBtn" onclick="run()">&#9654; Check text</button>
      <button class="sec" onclick="clearAll()">Clear</button>
      <span id="err" class="err"></span>
    </div>
  </div>

  <div class="card" id="libBuilder"><h2>1B &middot; Library builder <span class="tag">upload documents &rarr; feed words</span></h2>
    <p class="note">Upload one or more documents. STELOS lists the words that are in <b>neither</b> the STE dictionary
    <b>nor</b> the shared library, with an example sentence. Tick the real technical words, choose a category
    and add them. Every addition is recorded (who, when, category) for audit. Words seen here also help the tool
    repair broken words (for example &ldquo;use d&rdquo;) correctly.</p>
    <div class="row">
      <input type="file" id="hFiles" multiple accept=".docx,.xlsx,.xlsm,.txt,.md">
      <label class="opt">Pages <input type="text" id="hPages" placeholder="all (e.g. 1-3,5)" style="width:120px;padding:7px 9px;border:1px solid var(--border);border-radius:6px;font-size:13px"></label>
      <button id="hBtn" onclick="harvest()">&#9654; Find new words</button>
      <span id="hMsg" class="err"></span>
    </div>
    <div id="hOut" class="hidden">
      <div class="audbar">
        <label class="opt" style="font-size:12px"><input type="checkbox" onclick="document.querySelectorAll('.hchk').forEach(c=>c.checked=this.checked)"> Tick all</label>
        <label class="opt" style="font-size:12px">Add ticked words to:
          <select id="hCat">{% for c in term_categories %}<option value="{{c}}">{{c}}</option>{% endfor %}</select></label>
        <button class="tech" onclick="harvestAdd()">&#43; Add to library</button>
        <span id="hAddMsg"></span>
      </div>
      <p class="note" id="hSummary"></p>
      <table><thead><tr><th class="chk"></th><th>Word</th><th>Count</th><th>Example sentence</th></tr></thead>
        <tbody id="hBody"></tbody></table>
    </div>
  </div>

  <div id="results" class="hidden">
    <div id="wbNotes"></div>
    <div class="card"><h2>2 &middot; Summary</h2><div class="stats" id="stats"></div></div>

    <div class="card"><h2>3 &middot; Not-approved words &amp; replacement</h2>
      <p class="note">Words with an STE alternative: tick the left box and adjust the replacement, then apply.
      Words marked <i>"Not in STE dictionary"</i> are usually your technical terms &mdash; tick
      <b>"Add as tech word"</b> and save them to the shared library so they are approved (green) for
      everyone next time. Capitalization is matched automatically.</p>
      <table><thead><tr>
        <th class="chk"><input type="checkbox" id="allChk" onclick="toggleAll(this)"></th>
        <th>Word in text</th><th>POS (Parts of Speech) / usage <a class="qlink" href="#posApp" onclick="openPos()" title="What is POS? Open the Parts of Speech reference">?</a></th><th>Count</th><th>Use instead (STE) &amp; guidance</th><th>Replace with / tech word</th>
      </tr></thead><tbody id="wordsBody"></tbody></table>
      <div class="row">
        <button class="go" id="applyBtn" onclick="applyRepl()">&#10003; Apply ticked replacements &amp; download copy</button>
        <label class="opt" style="font-size:12px">Add ticked words to:
          <select id="termCat" style="padding:6px 8px;border:1px solid var(--border);border-radius:6px;font-size:12px">
            {% for c in term_categories %}<option value="{{c}}">{{c}}</option>{% endfor %}
          </select>
        </label>
        <button class="tech" id="techBtn" onclick="addTechWords()">&#43; Add to library</button>
        <span id="applyMsg"></span>
      </div>
      <div class="lib">
        <h4>Shared technical-terms library <span id="libUser" class="tag"></span></h4>
        <p class="note">Saved to a shared file, visible to everyone, grouped by category for maintenance and auditing.
        These words are treated as approved (green) and are never flagged.</p>
        <div id="libChips"></div>
      </div>
    </div>

    <div class="card" id="ctxCard"><h2>3A &middot; Cleared by sentence context <span class="tag" id="ctxCount"></span></h2>
      <p class="note">These words are <b>not</b> flagged because the sentence shows they are used in a part of speech
      in which they are approved (Rule 1.2) &mdash; for example &ldquo;test&rdquo; after &ldquo;the&rdquo; is a noun.
      This uses fixed rules, not a language model, so please check the list. If one is wrong, tell the team so the
      cue words in <code>context_rules.json</code> can be corrected.</p>
      <div id="ctxList"></div></div>

    <div class="card"><h2>4 &middot; Rule findings (Part 1 writing rules)</h2>
      <p class="note">Each finding shows the ASD-STE100 rule number. Errors break a rule outright;
      warnings and advisories need a quick check.</p>
      <div id="issues"></div></div>

    <div class="card"><h2>5 &middot; ASD-STE100 rule coverage</h2>
      <p class="note">Every rule in the standard. <b>Auto</b> = this tool checks it (issue count shown);
      <b>Manual</b> = needs your judgement (semantic rules a tool cannot decide).</p>
      <div id="coverage"></div></div>

    <div class="card"><h2>6 &middot; Annotated text</h2>
      <p class="note"><mark class="bad">red</mark> = not approved (hover for alternative) &nbsp;
      <mark class="warnw">amber</mark> = not in dictionary (strict mode)</p>
      <div class="doc" id="doc"></div>
      <div class="row">
        <button class="sec" onclick="dl()">&#8681; Download report (.html)</button>
        <button class="sec" onclick="annotateDoc()" id="annBtn">&#8681; Download annotated document (same format as input)</button>
        <span id="annMsg"></span>
      </div>
      <p class="note">The annotated document is a copy of your original file, text unchanged, with a reviewable
      comment added on each flagged occurrence (Word: native comment on that word, right-click to resolve/delete
      per occurrence; Excel: cell comment; text: inline [STE: ...]). Your original is never changed.</p>
    </div>
  </div>

  <div class="card" id="libAudit"><h2>Shared library &mdash; audit &amp; maintenance <span class="tag" id="audCount"></span></h2>
    <p class="note">One shared file for everybody. Filter by category, move words between categories, remove wrong
    entries, export the list to CSV, and see the change history (who added, moved or removed what).</p>
    <div class="audbar">
      <select id="audFilter" onchange="renderAudit()"><option value="">All categories</option>
        {% for c in term_categories %}<option value="{{c}}">{{c}}</option>{% endfor %}</select>
      <input type="text" id="audSearch" placeholder="Search word or user" oninput="renderAudit()">
      <label class="opt" style="font-size:12px">Move ticked to:
        <select id="audMoveCat">{% for c in term_categories %}<option value="{{c}}">{{c}}</option>{% endfor %}</select></label>
      <button class="sec" onclick="audMove()">Move</button>
      <button class="sec" onclick="audRemove()">Remove ticked</button>
      <a class="sec" href="/terms_export" style="text-decoration:none;padding:7px 12px;border:1px solid var(--border);border-radius:6px;font-size:12px">&#8681; Export CSV</a>
      <button class="sec" onclick="audHistory()">Change history</button>
      <span id="audMsg"></span>
    </div>
    <div id="audChips"></div>
    <div id="audHist" class="hidden"></div>
  </div>

  <details class="card" id="posApp"><summary style="cursor:pointer;font-weight:700;font-size:15px">Appendix &middot; Parts of Speech (POS) reference</summary>
    <p class="note" id="posStatus" style="margin-top:8px"></p>
    <p class="note" id="posIntro"></p>
    <div style="overflow-x:auto"><table class="postbl"><thead><tr><th>POS</th><th>Name</th><th>Definition</th><th>How it matters in STE</th><th>Example</th></tr></thead>
      <tbody id="posBody"><tr><td colspan="5">Loading...</td></tr></tbody></table></div>
    <p class="note">In the results, <b>POS</b> means <b>Parts of Speech</b>. The letters (N, V, ADJ ...) are the
    abbreviations used in the ASD-STE100 dictionary. The Linguistic Team can edit <code>pos_definitions.json</code> to
    change this text; no code change is needed.</p>
  </details>
</main>
<footer>STELOS &mdash; The ASD-STE100 Intelligence Platform &nbsp;&middot;&nbsp; &copy; 2026 ALTEN. All rights reserved. Confidential &ndash; Internal Use Only.</footer>
<script>
let LAST=null, TOKEN=null;
const STE_META={source:"{{ste_source}}", issue:"{{ste_issue}}", issue_date:"{{ste_issue_date}}", url:"{{ste_doc_url}}"};
function clearAll(){document.getElementById('text').value='';document.getElementById('file').value='';
document.getElementById('results').classList.add('hidden');document.getElementById('err').textContent='';
document.getElementById('pageRow').style.display='none';document.getElementById('pages').value='';}

async function onFilePick(){
  const f=document.getElementById('file').files[0];
  const row=document.getElementById('pageRow'), info=document.getElementById('pageInfo');
  document.getElementById('pages').value='';
  if(!f){row.style.display='none';return;}
  info.textContent='Reading pages...';row.style.display='flex';
  try{
    const fd=new FormData();fd.append('file',f);
    const r=await fetch('/pages',{method:'POST',body:fd});const d=await r.json();
    if(!r.ok){row.style.display='none';return;}
    const unit=d.unit||'page';
    info.innerHTML='Document has <b>'+d.count+' '+unit+(d.count>1?'s':'')+
      '</b>. Leave blank to check all, or enter e.g. <code>1-3,5</code>.';
  }catch(e){row.style.display='none';}
}
function esc(x){const e=document.createElement('div');e.textContent=x==null?'':x;return e.innerHTML;}
function escMark(text, word){
  const e=esc(text);
  const re=new RegExp('('+esc(word).replace(/[.*+?^${}()|[\]\\]/g,'\\$&')+')','gi');
  return e.replace(re,'<mark class="bad">$1</mark>');
}
function toggleParas(i, btn){
  const row=document.getElementById('paras-'+i);
  if(!row) return;
  const nowOpen=row.classList.toggle('hidden')===false;
  if(btn) btn.innerHTML=btn.innerHTML.replace(nowOpen?'▸':'▾', nowOpen?'▾':'▸');
}

async function run(){
  const err=document.getElementById('err');err.textContent='';
  const btn=document.getElementById('runBtn');btn.disabled=true;btn.textContent='Checking...';
  const fd=new FormData();
  fd.append('text',document.getElementById('text').value);
  fd.append('mode',document.getElementById('mode').value);
  fd.append('strict',document.getElementById('strict').checked?'true':'false');
  fd.append('terms',document.getElementById('terms').value);
  fd.append('pages',document.getElementById('pages').value);
  const f=document.getElementById('file').files[0];if(f)fd.append('file',f);
  try{
    const r=await fetch('/check',{method:'POST',body:fd});
    const d=await r.json();
    if(!r.ok){err.textContent=d.error||'Error';return;}
    LAST=d;TOKEN=d.token;render(d);
  }catch(e){err.textContent=e.message;}
  finally{btn.disabled=false;btn.innerHTML='&#9654; Check text';}
}

function render(d){
  document.getElementById('results').classList.remove('hidden');
  document.getElementById('applyMsg').textContent='';
  const s=d.summary;
  document.getElementById('stats').innerHTML=`
    <div class="stat"><div class="n">${s.words}</div><div class="l">Words</div></div>
    <div class="stat"><div class="n">${s.sentences}</div><div class="l">Sentences</div></div>
    <div class="stat ${s.unapproved_total?'bad':'ok'}"><div class="n">${s.unapproved_total}</div><div class="l">Not-approved uses</div></div>
    <div class="stat warn"><div class="n">${s.length_issues}</div><div class="l">Long sentences</div></div>
    <div class="stat warn"><div class="n">${s.paragraph_issues}</div><div class="l">Long paragraphs</div></div>
    <div class="stat ${s.rule_issues?'warn':'ok'}"><div class="n">${s.rule_issues}</div><div class="l">Rule findings</div></div>`;
  const wb=document.getElementById('wordsBody');
  if(d.flagged.length===0){wb.innerHTML='<tr><td colspan="6" style="color:var(--green)">No not-approved words found. &#10003;</td></tr>';}
  else{wb.innerHTML=d.flagged.map((f,i)=>{
     let lastcell, leftbox;
     if(f.kind==='unknown'){
       leftbox='';
       lastcell=`<td class="techcell"><label><input type="checkbox" class="techchk" data-word="${esc(f.word)}">
                  &#43; Add as tech word</label></td>`;
     } else if(f.reword || !f.has_alt){
       // verb->noun or no-direct-alternative: rewording, not a word swap
       leftbox='';
       lastcell=`<td class="repcell"><span class="rewordtag">Reword needed</span></td>`;
     } else {
       leftbox=`<input type="checkbox" class="rowchk" data-i="${i}" ${f.replacement?'checked':''} onchange="syncParas(${i},this.checked)">`;
       const multi=(f.alts&&f.alts.length>1);
       const sel = multi
         ? `<select class="repsel" data-i="${i}" onchange="onSel(${i})">
              ${f.alts.map(a=>`<option value="${esc(a.toLowerCase())}">${esc(a)}</option>`).join('')}
            </select> ` : '';
       lastcell=`<td class="repcell">${sel}<input class="rep" data-i="${i}"
                  value="${esc(f.replacement)}" placeholder="type replacement / rewrite..."></td>`;
     }
     const canRestrict = f.kind !== 'unknown';
     const occ = canRestrict ? (f.occurrences||[]) : [];
     const wordCell = occ.length
       ? `<button type="button" class="paratoggle" onclick="toggleParas(${i},this)">&#9656; ${occ.length} para${occ.length>1?'s':''}</button> ${esc(f.word)}`
       : esc(f.word);
     const paraRows = occ.length ? `<tr id="paras-${i}" class="hidden paralistrow"><td></td>
       <td colspan="5"><div class="paralist">${occ.map(o=>
         `<div class="paraitem${f.replacement?'':' unchecked'}"><label class="paraitemchk"><input type="checkbox" class="parachk" data-i="${i}" data-para="${o.para-1}" ${f.replacement?'checked':''}
             onchange="onParaTick(${i},this)"></label><b>Para ${o.para}:</b> ${escMark(o.text, f.word)}</div>`).join('')}
       </div></td></tr>` : '';
     const posCell = f.pos_display ? `<span class="posbadge">${esc(f.pos_display)}</span>` : '<span class="note" style="margin:0">&mdash;</span>';
     const useCell = `<div style="color:var(--green);font-family:var(--mono)">${esc(f.alt)}</div>`
        + (f.guidance ? `<div class="guidance">${esc(f.guidance)}</div>` : '');
     return `<tr>
       <td class="chk">${leftbox}</td>
       <td class="w">${wordCell}</td><td>${posCell}</td><td>${f.count}</td>
       <td>${useCell}</td>
       ${lastcell}</tr>${paraRows}`;}).join('');}
  renderLibrary(d.library, d.user);
  const wn=d.word_break_notes||[];
  document.getElementById('wbNotes').innerHTML=wn.length?`<div class="wbnote"><b>Possible broken words (${wn.length}):</b>
    ${wn.slice(0,12).map(x=>'<code>'+esc(x)+'</code>').join(', ')}${wn.length>12?' ...':''}<br>
    These look like one word split by a space (for example &ldquo;mix ed&rdquo;) but the joined word is not in the
    dictionary or library, so the text was left as it is. Check the source document.</div>`:'';
  const cc=d.context_cleared||[];
  document.getElementById('ctxCount').textContent=cc.length+' word(s)';
  document.getElementById('ctxList').innerHTML=cc.length?cc.slice(0,100).map(c=>
    `<div class="ctxrow"><b>${esc(c.word)}</b> <span class="why">${esc(c.reason)}</span>
     <div class="snip">${escMark(c.context,c.word)}</div></div>`).join('')
    +(cc.length>100?'<p class="note">Showing the first 100.</p>':''):'<p class="note">Nothing was cleared by context in this text.</p>';
  applyPilot();
  const iss=document.getElementById('issues');
  const rf=d.rule_findings||[];
  if(rf.length===0){iss.innerHTML='<p style="color:var(--green)">No rule findings. &#10003;</p>';}
  else{
    const order={error:0,warn:1,advisory:2};
    const sevlabel={error:'ERROR',warn:'WARNING',advisory:'ADVISORY'};
    const sorted=rf.slice().sort((a,b)=>(order[a.severity]-order[b.severity]));
    iss.innerHTML=sorted.map(i=>`<div class="issue sev-${i.severity}">
       <span class="rulebadge">Rule ${esc(i.rule)}</span>
       <span class="sevtag ${i.severity}">${sevlabel[i.severity]||''}</span>
       ${esc(i.msg)}<div class="ctx">${esc(i.context)}</div></div>`).join('');
  }
  const cov=document.getElementById('coverage');
  cov.innerHTML=(d.coverage||[]).map(sec=>`
    <div class="covsec"><h4>${esc(sec.section)}</h4>
    <div class="covgrid">${sec.rules.map(r=>{
      const auto=r.kind==='auto';
      const cls=auto?(r.count?'cov-hit':'cov-ok'):'cov-man';
      const tag=auto?(r.count?(r.count+' issue'+(r.count>1?'s':'')):'clear'):'manual';
      return `<div class="covitem ${cls}"><b>${esc(r.num)}</b> ${esc(r.title)}
              <span class="covtag">${tag}</span></div>`;}).join('')}</div></div>`).join('');
  document.getElementById('doc').innerHTML=d.highlighted;
}

function syncParas(i, on){
  // main word box -> all paragraph boxes under it follow
  document.querySelectorAll('.parachk[data-i="'+i+'"]').forEach(c=>{
    c.checked=on; c.closest('.paraitem').classList.toggle('unchecked', !on);});
}
function onParaTick(i, cb){
  cb.closest('.paraitem').classList.toggle('unchecked', !cb.checked);
  // paragraph box -> main word box is ticked while at least one paragraph is ticked
  const any=[...document.querySelectorAll('.parachk[data-i="'+i+'"]')].some(c=>c.checked);
  const main=document.querySelector('.rowchk[data-i="'+i+'"]');
  if(main) main.checked=any;
}
function toggleAll(cb){document.querySelectorAll('.rowchk').forEach(c=>{c.checked=cb.checked;syncParas(c.dataset.i,cb.checked);});}
function onSel(i){
  const sel=document.querySelector('.repsel[data-i="'+i+'"]');
  const inp=document.querySelector('.rep[data-i="'+i+'"]');
  if(sel&&inp) inp.value=sel.value;
  const chk=document.querySelector('.rowchk[data-i="'+i+'"]');
  if(chk){chk.checked=true;syncParas(i,true);}
}

let LIB={};
function renderLibrary(lib, user){
  LIB=lib||{};
  try{renderAudit();}catch(e){}
  document.getElementById('libUser').textContent = user ? ('you: '+user) : '';
  const box=document.getElementById('libChips');
  const keys=Object.keys(lib||{}).sort();
  if(keys.length===0){box.innerHTML='<span class="note">No technical terms saved yet.</span>';return;}
  // item 5 - group the library by category for maintenance / auditing
  const groups={};
  keys.forEach(k=>{const cat=(lib[k]&&lib[k].category)||'Miscellaneous';(groups[cat]=groups[cat]||[]).push(k);});
  const order=['Technical Terms','Material Names','Part Names','Miscellaneous'];
  const cats=Object.keys(groups).sort((a,b)=>{
    const ia=order.indexOf(a),ib=order.indexOf(b);return (ia<0?99:ia)-(ib<0?99:ib);});
  box.innerHTML=cats.map(cat=>`<div class="libcat"><h5>${esc(cat)} (${groups[cat].length})</h5>
    <div class="chips">${groups[cat].map(k=>
      `<span class="chip">${esc(k)} <small>&middot; ${esc((lib[k]&&lib[k].by)||'')}</small></span>`).join('')}</div></div>`).join('');
}

// ---- library audit & maintenance -------------------------------------------
function renderAudit(){
  const box=document.getElementById('audChips');if(!box)return;
  const cat=document.getElementById('audFilter').value;
  const q=(document.getElementById('audSearch').value||'').trim().toLowerCase();
  const keys=Object.keys(LIB).sort().filter(k=>{
    const v=LIB[k]||{};
    if(cat && (v.category||'Miscellaneous')!==cat) return false;
    if(q && k.indexOf(q)<0 && String(v.by||'').toLowerCase().indexOf(q)<0) return false;
    return true;});
  document.getElementById('audCount').textContent=keys.length+' of '+Object.keys(LIB).length+' words';
  if(!keys.length){box.innerHTML='<span class="note">No words match.</span>';return;}
  const groups={};
  keys.forEach(k=>{const c=(LIB[k]&&LIB[k].category)||'Miscellaneous';(groups[c]=groups[c]||[]).push(k);});
  const order=['Technical Terms','Material Names','Part Names','Miscellaneous'];
  const cats=Object.keys(groups).sort((a,b)=>{const ia=order.indexOf(a),ib=order.indexOf(b);return (ia<0?99:ia)-(ib<0?99:ib);});
  box.innerHTML=cats.map(c=>`<div class="libcat"><h5>${esc(c)} (${groups[c].length})</h5>
    <div class="chips">${groups[c].map(k=>{const v=LIB[k]||{};
      return `<label class="chip" title="Added by ${esc(v.by||'')} on ${esc(v.at||'')}"><input type="checkbox" class="audchk" data-word="${esc(k)}">${esc(k)}
        <small>&middot; ${esc(v.by||'')} ${esc((v.at||'').slice(0,10))}</small></label>`;}).join('')}</div></div>`).join('');
}
function _audWords(){return [...document.querySelectorAll('.audchk')].filter(c=>c.checked).map(c=>c.dataset.word);}
async function _audPost(url, body, okText){
  const msg=document.getElementById('audMsg');msg.className='';msg.textContent='';
  try{
    const r=await fetch(url,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(body)});
    const d=await r.json();
    if(!r.ok){msg.className='err';msg.textContent=d.error||'Error';return;}
    renderLibrary(d.library,d.user);
    msg.className='ok-msg';msg.textContent=okText(d);
  }catch(e){msg.className='err';msg.textContent=e.message;}
}
function audMove(){
  const w=_audWords();const msg=document.getElementById('audMsg');
  if(!w.length){msg.className='err';msg.textContent='Tick at least one word.';return;}
  const cat=document.getElementById('audMoveCat').value;
  _audPost('/move_term',{terms:w,category:cat},d=>'Moved '+d.moved.length+' word(s) to "'+cat+'".');
}
function audRemove(){
  const w=_audWords();const msg=document.getElementById('audMsg');
  if(!w.length){msg.className='err';msg.textContent='Tick at least one word.';return;}
  if(!confirm('Remove '+w.length+' word(s) from the shared library for everyone?'))return;
  _audPost('/remove_term',{terms:w},d=>'Removed '+d.removed.length+' word(s).');
}
async function audHistory(){
  const box=document.getElementById('audHist');
  if(box.classList.toggle('hidden')){return;}
  box.innerHTML='Loading...';
  try{
    const r=await fetch('/terms_history');const d=await r.json();
    const h=d.history||[];
    box.innerHTML=h.length?('<table><thead><tr><th>When</th><th>User</th><th>Action</th><th>Word</th><th>From</th><th>To</th></tr></thead><tbody>'
      +h.map(x=>`<tr class="hrow"><td>${esc(x['Timestamp'])}</td><td>${esc(x['User'])}</td><td>${esc(x['Action'])}</td><td>${esc(x['Word'])}</td><td>${esc(x['From category'])}</td><td>${esc(x['To category'])}</td></tr>`).join('')
      +'</tbody></table>'):'<p class="note">No changes recorded yet.</p>';
  }catch(e){box.textContent=e.message;}
}

// ---- library builder (upload documents -> feed words) ------------------------
async function harvest(){
  const msg=document.getElementById('hMsg');msg.textContent='';
  const files=document.getElementById('hFiles').files;
  if(!files.length){msg.textContent='Choose one or more documents.';return;}
  const btn=document.getElementById('hBtn');btn.disabled=true;btn.textContent='Reading...';
  const fd=new FormData();
  for(const f of files) fd.append('files',f);
  fd.append('pages',document.getElementById('hPages').value);
  try{
    const r=await fetch('/harvest',{method:'POST',body:fd});const d=await r.json();
    if(!r.ok){msg.textContent=d.error||'Error';return;}
    document.getElementById('hOut').classList.remove('hidden');
    document.getElementById('hSummary').textContent=d.documents.length+' document(s), '+d.words+' words read. '+
      d.total_unknown+' new word(s) found'+(d.total_unknown>d.rows.length?' (showing the '+d.rows.length+' most frequent)':'')+'.';
    document.getElementById('hBody').innerHTML=d.rows.length?d.rows.map(x=>`<tr>
      <td class="chk"><input type="checkbox" class="hchk" data-word="${esc(x.word)}"></td>
      <td class="w">${esc(x.word)}${x.capitalised?' <span class="note" style="margin:0">(capitalised)</span>':''}</td>
      <td>${x.count}</td><td style="font-size:12px">${escMark(x.example,x.word)}</td></tr>`).join('')
      :'<tr><td colspan="4" style="color:var(--green)">No new words - everything is already in the dictionary or library.</td></tr>';
  }catch(e){msg.textContent=e.message;}
  finally{btn.disabled=false;btn.innerHTML='&#9654; Find new words';}
}
async function harvestAdd(){
  const msg=document.getElementById('hAddMsg');msg.className='';msg.textContent='';
  const words=[...document.querySelectorAll('.hchk')].filter(c=>c.checked).map(c=>c.dataset.word);
  if(!words.length){msg.className='err';msg.textContent='Tick at least one word.';return;}
  const cat=document.getElementById('hCat').value;
  try{
    const r=await fetch('/add_term',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({terms:words,category:cat})});
    const d=await r.json();
    if(!r.ok){msg.className='err';msg.textContent=d.error||'Error';return;}
    renderLibrary(d.library,d.user);
    document.querySelectorAll('.hchk').forEach(c=>{if(c.checked){c.closest('tr').remove();}});
    msg.className='ok-msg';msg.textContent='Added '+d.added.length+' word(s) to "'+cat+'".';
  }catch(e){msg.className='err';msg.textContent=e.message;}
}

// ---- Parts of Speech (POS) appendix -------------------------------------------
function openPos(){const p=document.getElementById('posApp');p.open=true;setTimeout(()=>p.scrollIntoView({behavior:'smooth'}),50);}
async function loadPos(){
  try{
    const r=await fetch('/pos_ref');const d=await r.json();
    document.getElementById('posStatus').innerHTML='<i>'+esc(d.status||'')+'</i>';
    document.getElementById('posIntro').textContent=d.intro||'';
    document.getElementById('posBody').innerHTML=(d.entries||[]).map(e=>`<tr>
      <td class="abbr">${esc(String(e.abbr||'').toUpperCase())}</td><td><b>${esc(e.name)}</b></td>
      <td>${esc(e.definition)}</td><td>${esc(e.ste_tip)}</td>
      <td>${e.ste?'<div class="stx"><b>STE:</b> '+esc(e.ste)+'</div>':''}${e.non_ste?'<div class="nstx"><b>Non-STE:</b> '+esc(e.non_ste)+'</div>':''}</td></tr>`).join('');
  }catch(e){document.getElementById('posBody').innerHTML='<tr><td colspan="5">Could not load the reference.</td></tr>';}
}

// ---- pilot mode ---------------------------------------------------------------
const PILOT={{ 'true' if pilot_mode else 'false' }};
function applyPilot(){
  if(!PILOT) return;
  [['applyBtn','Apply replacements'],['annBtn','Annotated document']].forEach(([id,label])=>{
    const b=document.getElementById(id);if(!b)return;
    b.disabled=true;b.title='Switched off in pilot mode';b.innerHTML=label+' (off - pilot mode)';});
}

async function loadLibrary(){
  try{const r=await fetch('/terms');const d=await r.json();renderLibrary(d.library,d.user);}catch(e){}
}

async function addTechWords(){
  const msg=document.getElementById('applyMsg');msg.className='';msg.textContent='';
  const words=[...document.querySelectorAll('.techchk')].filter(c=>c.checked).map(c=>c.dataset.word);
  if(words.length===0){msg.className='err';msg.textContent='Tick at least one "Add as tech word" box.';return;}
  const btn=document.getElementById('techBtn');btn.disabled=true;btn.textContent='Saving...';
  try{
    const cat=document.getElementById('termCat').value;
    const r=await fetch('/add_term',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({terms:words, category:cat})});
    const d=await r.json();
    if(!r.ok){msg.className='err';msg.textContent=d.error||'Error';return;}
    renderLibrary(d.library,d.user);
    msg.className='ok-msg';msg.textContent='Saved '+d.added.length+' term(s) to "'+cat+'". Re-checking...';
    // grey the added rows immediately, then re-check so they turn green / drop off
    document.querySelectorAll('.techchk').forEach(c=>{if(c.checked)c.closest('tr').classList.add('rowdone');});
    setTimeout(run, 500);
  }catch(e){msg.className='err';msg.textContent=e.message;}
  finally{btn.disabled=false;btn.innerHTML='&#43; Add to library';}
}

async function applyRepl(){
  const msg=document.getElementById('applyMsg');msg.className='';msg.textContent='';
  if(!TOKEN){msg.className='err';msg.textContent='Run a check first.';return;}
  const reps=[];
  document.querySelectorAll('.rowchk').forEach(chk=>{
    if(chk.checked){
      const i=chk.dataset.i;const f=LAST.flagged[i];
      const rep=document.querySelector('.rep[data-i="'+i+'"]').value.trim();
      if(!rep) return;
      const entry={find:f.word,replace:rep};
      // Paragraph indices are relative to the checked text; if a page filter
      // was used, those indices don't map to the full document, so apply the
      // ticked word across the whole document instead of per-paragraph.
      const pageFilter=document.getElementById('pages').value.trim();
      const occChecks=[...document.querySelectorAll('.parachk[data-i="'+i+'"]')];
      if(occChecks.length && !pageFilter){
        const chosen=occChecks.filter(c=>c.checked).map(c=>parseInt(c.dataset.para,10));
        if(chosen.length<occChecks.length) entry.paras=chosen;  // restrict to ticked paragraphs only
      }
      reps.push(entry);
    }
  });
  if(reps.length===0){msg.className='err';msg.textContent='Tick at least one word that has a replacement.';return;}
  const btn=document.getElementById('applyBtn');btn.disabled=true;btn.textContent='Working...';
  try{
    const r=await fetch('/apply',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({token:TOKEN,replacements:reps})});
    if(!r.ok){const d=await r.json();msg.className='err';msg.textContent=d.error||'Error';return;}
    const n=r.headers.get('X-STE-Replacements');const name=r.headers.get('X-STE-Outfile')||'corrected';
    const blob=await r.blob();const a=document.createElement('a');
    a.href=URL.createObjectURL(blob);a.download=name;a.click();
    msg.className='ok-msg';msg.textContent='Done - '+n+' replacement(s) written to a copy: '+name;
  }catch(e){msg.className='err';msg.textContent=e.message;}
  finally{btn.disabled=false;btn.innerHTML='&#10003; Apply ticked replacements &amp; download copy';}
}

async function annotateDoc(){
  const msg=document.getElementById('annMsg');msg.className='';msg.textContent='';
  if(!TOKEN){msg.className='err';msg.textContent='Run a check first.';return;}
  const btn=document.getElementById('annBtn');btn.disabled=true;btn.textContent='Building...';
  try{
    const r=await fetch('/annotate',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({token:TOKEN,strict:document.getElementById('strict').checked,
        terms:document.getElementById('terms').value})});
    if(!r.ok){const d=await r.json();msg.className='err';msg.textContent=d.error||'Error';return;}
    const nm=r.headers.get('X-STE-Outfile')||'annotated';const nn=r.headers.get('X-STE-Annotations');
    const cmode=r.headers.get('X-STE-Comment-Mode');
    const blob=await r.blob();const a=document.createElement('a');
    a.href=URL.createObjectURL(blob);a.download=nm;a.click();
    if(cmode==='fallback'){
      msg.className='err';
      msg.textContent='Downloaded '+nm+' with '+nn+' mark(s), but this Python has an old python-docx '+
        '(no native Word comments) so words are highlighted with an inline [STE: ...] note instead of a real '+
        'comment. Run: pip install --upgrade python-docx  to get real comments.';
    } else if(cmode==='com'){
      msg.className='ok-msg';
      msg.textContent='Annotated document with '+nn+' real Word comment(s) downloaded: '+nm+
        ' (added via Word automation, since python-docx here is too old for native comments).';
    } else {
      msg.className='ok-msg';msg.textContent='Annotated document with '+nn+' mark(s) downloaded: '+nm;
    }
  }catch(e){msg.className='err';msg.textContent=e.message;}
  finally{btn.disabled=false;btn.innerHTML='&#8681; Download annotated document (same format as input)';}
}

function dl(){
  if(!LAST)return;
  const doc=document.getElementById('doc').innerHTML;const s=LAST.summary;
  const rows=LAST.flagged.map(f=>`<tr><td>${esc(f.word)}</td><td>${esc(f.pos_display||f.pos||'')}</td><td>${f.count}</td><td>${esc(f.alt)}${f.guidance?('<br><small style="color:#475569">'+esc(f.guidance)+'</small>'):''}</td></tr>`).join('');
  const isr=(LAST.rule_findings||[]).map(i=>`<li><b>Rule ${esc(i.rule)}</b> [${esc(i.severity)}] ${esc(i.msg)} <i>${esc(i.context)}</i></li>`).join('');
  const cov=(LAST.coverage||[]).map(sec=>`<h3>${esc(sec.section)}</h3><table><tr><th>Rule</th><th>Title</th><th>Coverage</th></tr>`+
    sec.rules.map(r=>`<tr><td>${esc(r.num)}</td><td>${esc(r.title)}</td><td>${r.kind==='auto'?(r.count?('AUTO - '+r.count+' issue(s)'):'AUTO - clear'):'Manual review'}</td></tr>`).join('')+`</table>`).join('');
  const h=`<!doctype html><meta charset=utf-8><title>STELOS - ASD-STE100 report</title>
  <link rel="icon" href="https://www.alten.com/wp-content/uploads/2019/01/favicon-alten.png">
  <style>body{font-family:Segoe UI,sans-serif;max-width:900px;margin:30px auto;color:#1e293b}
  h1{color:#1a4fad}h3{color:#475569;margin-top:14px}mark.bad{background:#fde2e2;border-bottom:2px solid #dc2626}mark.warnw{background:#fef3c7}
  table{border-collapse:collapse;width:100%;margin-bottom:8px}td,th{border:1px solid #c8d4e8;padding:6px;text-align:left;font-size:13px}
  .doc{background:#f1f5fb;padding:14px;border-radius:8px;line-height:1.7}</style>
  <h1>STELOS &ndash; ASD-STE100 report</h1><p style="color:#475569;margin-top:-6px">The ASD-STE100 Intelligence Platform &middot; An ALTEN product</p>
  <p><b>Validated against:</b> ${esc(STE_META.source)}${STE_META.url?(' &middot; <a href="'+esc(STE_META.url)+'">specification</a>'):''}</p>
  <p>${new Date().toLocaleString()} &middot; ${s.mode} mode (limit ${s.sentence_limit} words)</p>
  <p><b>${s.words}</b> words, <b>${s.sentences}</b> sentences, <b>${s.unapproved_total}</b> not-approved uses,
  <b>${s.length_issues}</b> long sentences, <b>${s.paragraph_issues}</b> long paragraphs, <b>${s.rule_issues}</b> rule findings.</p>
  <h2>Not-approved words</h2><table><tr><th>Word / rule</th><th>Tag</th><th>Count</th><th>Use instead</th></tr>${rows}</table>
  <h2>Rule findings</h2><ul>${isr||'<li>None</li>'}</ul>
  <h2>ASD-STE100 rule coverage</h2>${cov}
  <h2>Annotated text</h2><div class="doc">${doc}</div>
  <hr><small>STELOS &middot; &copy; 2026 ALTEN. Confidential - Internal Use Only.</small>`;
  const b=new Blob([h],{type:'text/html'});const a=document.createElement('a');
  a.href=URL.createObjectURL(b);a.download='STE_report_'+Date.now()+'.html';a.click();
}
loadLibrary();
loadPos();
applyPilot();
</script>
</body></html>"""

def open_browser():
    time.sleep(1.2)
    webbrowser.open('http://127.0.0.1:%d' % PORT)

if __name__ == '__main__':
    print('STELOS - The ASD-STE100 Intelligence Platform (ALTEN) | %s | %d approved / %d not-approved' %
          (META['source'], META['approved_count'], META['not_approved_count']))
    threading.Thread(target=open_browser, daemon=True).start()
    app.run(host='127.0.0.1', port=PORT, debug=False, use_reloader=False)
